import os
import sys
import django
from pathlib import Path
from openpyxl import load_workbook
from django.utils import timezone


BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.append(str(BASE_DIR))

os.environ.setdefault(
    "DJANGO_SETTINGS_MODULE",
    "config.settings"
)

django.setup()


from flota.models import Vehiculo, CentroCosto


RUTA_EXCEL = "scripts/base_admin.xlsx"


def normalizar_patente(valor):

    if valor is None:
        return None

    return (
        str(valor)
        .upper()
        .replace("-", "")
        .replace(" ", "")
        .strip()
    )


def limpiar_texto(valor):

    if valor is None:
        return None

    texto = str(valor).strip()

    if texto == "":
        return None

    return texto


def obtener_indice(encabezados, nombre_columna):

    if nombre_columna not in encabezados:
        return None

    return encabezados.index(nombre_columna)


wb = load_workbook(RUTA_EXCEL)
hoja = wb.active


encabezados = []

for celda in hoja[1]:

    encabezados.append(
        str(celda.value).strip()
        if celda.value is not None
        else ""
    )


indice_patente = obtener_indice(
    encabezados,
    "Patente"
)

if indice_patente is None:
    print("ERROR: No se encontró columna Patente.")
    exit()


indices = {
    "centro_costo": obtener_indice(encabezados, "Centro De Costo"),
    "empresa_arrendadora": obtener_indice(encabezados, "Empresa Arrendadora"),
    "razon_social": obtener_indice(encabezados, "Razón Social"),
    "rut_empresa": obtener_indice(encabezados, "Rut"),
    "sucursal": obtener_indice(encabezados, "Sucursal"),
    "ciudad": obtener_indice(encabezados, "Ciudad"),
}


datos_admin = {}
duplicados = {}


for fila in hoja.iter_rows(
    min_row=2,
    values_only=True
):

    patente_normalizada = normalizar_patente(
        fila[indice_patente]
    )

    if not patente_normalizada:
        continue

    if patente_normalizada in datos_admin:

        duplicados[patente_normalizada] = (
            duplicados.get(patente_normalizada, 1) + 1
        )

    def valor(campo):

        indice = indices.get(campo)

        if indice is None:
            return None

        return limpiar_texto(
            fila[indice]
        )

    datos_admin[patente_normalizada] = {
        "centro_costo": valor("centro_costo"),
        "empresa_arrendadora": valor("empresa_arrendadora"),
        "razon_social": valor("razon_social"),
        "rut_empresa": valor("rut_empresa"),
        "sucursal": valor("sucursal"),
        "ciudad": valor("ciudad"),
    }


patentes_admin = set(
    datos_admin.keys()
)


vehiculos_bd = Vehiculo.objects.all()

actualizados = 0
creados = 0
dados_baja = 0
reactivados = 0


for patente_admin, datos in datos_admin.items():

    vehiculo = Vehiculo.objects.filter(
        patente=patente_admin
    ).first()

    centro_costo = None

    if datos["centro_costo"]:

        centro_costo, _ = CentroCosto.objects.get_or_create(
            codigo=datos["centro_costo"],
            defaults={
                "nombre": datos["centro_costo"]
            }
        )

    if vehiculo:

        estaba_baja = (
            vehiculo.estado_administrativo == "Dado de Baja"
        )

        vehiculo.centro_costo = centro_costo
        vehiculo.empresa_arrendadora = datos["empresa_arrendadora"]
        vehiculo.razon_social = datos["razon_social"]
        vehiculo.rut_empresa = datos["rut_empresa"]
        vehiculo.sucursal = datos["sucursal"]
        vehiculo.ciudad = datos["ciudad"]

        vehiculo.estado_administrativo = "Vigente"
        vehiculo.fecha_baja = None

        if estaba_baja:
            reactivados += 1

        vehiculo.save()

        actualizados += 1

    else:

        Vehiculo.objects.create(
            patente=patente_admin,
            centro_costo=centro_costo,
            empresa_arrendadora=datos["empresa_arrendadora"],
            razon_social=datos["razon_social"],
            rut_empresa=datos["rut_empresa"],
            sucursal=datos["sucursal"],
            ciudad=datos["ciudad"],
            estado_administrativo="Vigente",
            estado_operacional="No informado",
        )

        creados += 1


for vehiculo in vehiculos_bd:

    patente_normalizada_bd = normalizar_patente(
        vehiculo.patente
    )

    if patente_normalizada_bd not in patentes_admin:

        if vehiculo.estado_administrativo != "Dado de Baja":

            vehiculo.estado_administrativo = "Dado de Baja"
            vehiculo.estado_operacional = "FueraServicio"
            vehiculo.fecha_baja = timezone.now().date()
            vehiculo.save()

            dados_baja += 1


print("")
print("=== CRUCE ADMIN APLICADO ===")
print("")

print("Patentes únicas admin:", len(patentes_admin))
print("Actualizados:", actualizados)
print("Creados:", creados)
print("Reactivados:", reactivados)
print("Dados de baja:", dados_baja)
print("Duplicados en Excel admin:", len(duplicados))

print("")
print("Proceso completado.")