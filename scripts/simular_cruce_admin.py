import os
import sys
import django
from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.append(str(BASE_DIR))


os.environ.setdefault(
    "DJANGO_SETTINGS_MODULE",
    "config.settings"
)

django.setup()


from flota.models import Vehiculo


RUTA_EXCEL = "scripts/base_admin.xlsx"


def normalizar_patente(valor):

    if valor is None:
        return None

    return (
        str(valor)
        .upper()
        .replace("-", "")
        .replace(" ", "")
        .strip()
    )


wb = load_workbook(RUTA_EXCEL)
hoja = wb.active


encabezados = []

for celda in hoja[1]:

    encabezados.append(
        str(celda.value).strip()
        if celda.value is not None
        else ""
    )


if "Patente" not in encabezados:

    print("ERROR: No se encontró la columna Patente.")
    exit()


indice_patente = encabezados.index("Patente")


patentes_admin = set()

filas_sin_patente = 0
total_filas_con_patente = 0
patentes_repetidas = {}

for fila in hoja.iter_rows(
    min_row=2,
    values_only=True
):

    patente = normalizar_patente(
        fila[indice_patente]
    )

    if patente:

        total_filas_con_patente += 1

        if patente in patentes_admin:

            patentes_repetidas[patente] = (
                patentes_repetidas.get(patente, 1) + 1
            )

        patentes_admin.add(patente)

    else:

        filas_sin_patente += 1


vehiculos_bd = Vehiculo.objects.all()

patentes_bd = set(
    vehiculos_bd.values_list(
        "patente",
        flat=True
    )
)


coinciden = patentes_bd.intersection(
    patentes_admin
)

solo_admin = patentes_admin.difference(
    patentes_bd
)

solo_bd = patentes_bd.difference(
    patentes_admin
)


print("")
print("=== SIMULACIÓN CRUCE ADMIN ===")
print("")

print(
    "Patentes en Excel admin:",
    len(patentes_admin)
)
print(
    "Filas con patente en Excel:",
    total_filas_con_patente
)

print(
    "Patentes únicas normalizadas:",
    len(patentes_admin)
)

print(
    "Patentes repetidas detectadas:",
    len(patentes_repetidas)
)

print(
    "Filas sin patente en Excel:",
    filas_sin_patente
)

print(
    "Vehículos en BD:",
    len(patentes_bd)
)

print("")
print(
    "Coinciden:",
    len(coinciden)
)

print(
    "Nuevos en Excel admin:",
    len(solo_admin)
)

print(
    "Se darían de baja:",
    len(solo_bd)
)

print("")

print("Ejemplos nuevos en Excel admin:")
for patente in list(solo_admin)[:10]:
    print("-", patente)

print("")

print("Ejemplos a dar de baja:")
for patente in list(solo_bd)[:10]:
    print("-", patente)

print("")
print("SIMULACIÓN COMPLETADA. No se modificó la BD.")

print("")
print("Ejemplos patentes repetidas:")
for patente, cantidad in list(patentes_repetidas.items())[:20]:
    print("-", patente, "aparece", cantidad, "veces")

