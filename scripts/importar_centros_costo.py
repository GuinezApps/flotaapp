import os
import sys
import django
from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.append(str(BASE_DIR))

os.environ.setdefault(
    "DJANGO_SETTINGS_MODULE",
    "config.settings"
)

django.setup()


from flota.models import CentroCosto


RUTA_EXCEL = BASE_DIR / "Centro de Costos Vigentes.xlsx"


def limpiar_codigo(valor):

    if valor is None:
        return None

    codigo = str(valor).strip()

    if codigo.endswith(".0"):
        codigo = codigo[:-2]

    if codigo == "":
        return None

    return codigo


def limpiar_texto(valor):

    if valor is None:
        return None

    texto = str(valor).strip()

    if texto == "":
        return None

    return texto


def main():

    if not RUTA_EXCEL.exists():
        print("")
        print("No se encontró el archivo:")
        print(RUTA_EXCEL)
        print("")
        print("Copia el Excel en la raíz del proyecto con este nombre:")
        print("Centro de Costos Vigentes.xlsx")
        return

    wb = load_workbook(
        RUTA_EXCEL,
        data_only=True
    )

    hoja = wb.active

    encabezados = {}

    for celda in hoja[2]:

        if celda.value:

            encabezados[
                str(celda.value).strip()
            ] = celda.column

    columnas_necesarias = [
        "CC",
        "NOMBRE CC",
        "ESTADO"
    ]

    faltantes = []

    for columna in columnas_necesarias:

        if columna not in encabezados:
            faltantes.append(
                columna
            )

    if faltantes:
        print("")
        print("Faltan columnas necesarias:")
        print(", ".join(faltantes))
        return

    col_cc = encabezados["CC"]
    col_nombre = encabezados["NOMBRE CC"]
    col_estado = encabezados["ESTADO"]

    creados = 0
    actualizados = 0
    sin_cambios = 0
    omitidos_no_vigentes = 0
    errores = []

    for fila in range(
        3,
        hoja.max_row + 1
    ):

        codigo = limpiar_codigo(
            hoja.cell(
                row=fila,
                column=col_cc
            ).value
        )

        nombre = limpiar_texto(
            hoja.cell(
                row=fila,
                column=col_nombre
            ).value
        )

        estado = limpiar_texto(
            hoja.cell(
                row=fila,
                column=col_estado
            ).value
        )

        if not codigo:
            continue

        if estado != "VIGENTE":
            omitidos_no_vigentes += 1
            continue

        if not nombre:
            errores.append(
                f"Fila {fila}: CC {codigo} sin nombre"
            )
            continue

        centro_costo, creado = CentroCosto.objects.get_or_create(
            codigo=codigo,
            defaults={
                "nombre": nombre
            }
        )

        if creado:
            creados += 1
            continue

        if centro_costo.nombre != nombre:

            centro_costo.nombre = nombre
            centro_costo.save()

            actualizados += 1

        else:

            sin_cambios += 1

    print("")
    print("=== IMPORTACIÓN CENTROS DE COSTO ===")
    print("")
    print(f"Creados: {creados}")
    print(f"Actualizados: {actualizados}")
    print(f"Sin cambios: {sin_cambios}")
    print(f"Omitidos no vigentes: {omitidos_no_vigentes}")
    print(f"Errores: {len(errores)}")

    if errores:
        print("")
        print("Detalle de errores:")
        for error in errores[:30]:
            print("-", error)

    print("")
    print("Proceso completado.")


if __name__ == "__main__":
    main()