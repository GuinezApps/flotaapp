from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from django.utils import timezone
from django.db.models import Q
from django.db.models.functions import Upper, Trim
from django.contrib import messages
from django.contrib.auth.models import User, Group
from django.contrib.auth.forms import SetPasswordForm
from django.contrib.auth.decorators import login_required
from datetime import timedelta

from openpyxl import Workbook
from openpyxl import load_workbook

from .models import (
    Vehiculo,
    PerfilUsuario,
    CentroCosto,
    BitacoraAccion,
    HistorialBajaVehiculo,
    HistorialTransferenciaVehiculo,
    MantencionVehiculo,
)

from .forms import (
    VehiculoForm,
    CrearUsuarioForm,
    MantencionVehiculoForm,
    CerrarMantencionVehiculoForm,
    CancelarMantencionVehiculoForm,
    ReprogramarMantencionVehiculoForm,
)

from .decorators import (
    editor_required,
    master_required,
)

from .auditoria import registrar_bitacora

EXCEL_MAP = {
    'Numero Interno De Faena': 'numero_interno_faena',
    'Patente': 'patente',
    'Año': 'anio',
    'Combustible': 'combustible',
    'Tracción': 'traccion',
    'Nro. Motor': 'numero_motor',
    'Nro. Chasis/VIN': 'numero_chasis_vin',
    'Tipo Vehículo': 'tipo_vehiculo',
    'Marca': 'marca',
    'Modelo': 'modelo',
    'Color': 'color',
    'Nro. Asientos': 'numero_asientos',
    'Tipo De Moneda': 'tipo_moneda',
    'Tarifa $/Mes': 'tarifa_mensual',
    'Centro De Costo': 'centro_costo',
    'Tipo Estándar': 'tipo_estandar',
    'Fecha Inicio Contrato': 'fecha_inicio_contrato',
    'Fecha Término Contrato': 'fecha_termino_contrato',
    'Período (Meses)': 'periodo_meses',
    'Límite Del Kilometraje': 'limite_kilometraje',
    'Kilometraje Última Mantención': 'kilometraje_ultima_mantencion',
    'Kilometraje Proxima Mantención': 'kilometraje_proxima_mantencion',
    'Kilometraje Estimado Diario': 'kilometraje_estimado_diario',
    'Fecha Estimada De Mantención': 'fecha_estimada_mantencion',
    'Kilometraje De La Camioneta': 'kilometraje_actual',
    'Kilometraje Restantes': 'kilometraje_restante',
    'Días Restantes': 'dias_restantes',
    'Alerta De Mantención': 'alerta_mantencion',
    'Empresa Arrendadora': 'empresa_arrendadora',
    'Razón Social': 'razon_social',
    'Rut': 'rut_empresa',
    'Sucursal': 'sucursal',
    'Ciudad': 'ciudad',
    'Faena': 'faena',
    'Gerencia': 'gerencia',
    'Fecha De Creacion': 'fecha_creacion',
    'Fecha En Que Se Dio De Baja': 'fecha_baja',
    'Estado': 'estado_mantencion',
}


def limpiar_texto(valor):
    if valor is None:
        return None

    texto = str(valor).strip()

    if texto == "" or texto == "-":
        return None

    return texto


def normalizar_marca_filtro(valor):
    if not valor:
        return ''

    texto = str(valor)

    texto = texto.replace(
        '\xa0',
        ' '
    )

    texto = texto.replace(
        '\u200b',
        ''
    )

    texto = texto.strip().upper()

    texto = ' '.join(
        texto.split()
    )

    return texto


def limpiar_int(valor):
    if valor is None or valor == "":
        return None

    try:
        return int(float(str(valor).replace(",", ".")))
    except ValueError:
        return None


def limpiar_fecha(valor):
    if valor is None or valor == "":
        return None

    if hasattr(valor, "date"):
        return valor.date()

    return None


def limpiar_tarifa(valor):
    if valor is None:
        return None

    return str(valor).strip()


def obtener_encabezados(hoja):
    encabezados = []

    # El Excel real trae encabezados en la fila 2.
    for celda in hoja[2]:
        encabezados.append(
            str(celda.value).strip()
            if celda.value is not None
            else ""
        )

    return encabezados


def obtener_datos_fila(encabezados, fila):
    datos_originales = dict(zip(encabezados, fila))
    datos = {}

    for columna_excel, campo_modelo in EXCEL_MAP.items():
        datos[campo_modelo] = datos_originales.get(columna_excel)

    return datos


def obtener_o_crear_centro_costo(codigo):
    codigo_limpio = limpiar_texto(codigo)

    if not codigo_limpio:
        return None

    centro_costo, _ = CentroCosto.objects.get_or_create(
        codigo=codigo_limpio,
        defaults={
            'nombre': codigo_limpio
        }
    )

    return centro_costo


def obtener_estado_administrativo(fecha_baja):
    if fecha_baja:
        return 'Dado de Baja'

    return 'Vigente'

def normalizar_estado_operacional(estado):
    conversiones = {
        'Operativo': ('Operativo', None),
        'En Mantencion': ('No operativo', 'Mantencion'),
        'En Mantención': ('No operativo', 'Mantencion'),
        'FueraServicio': ('No operativo', 'Detenido'),
        'Fuera de Servicio': ('No operativo', 'Detenido'),
        'No operativo': ('No operativo', None),
        'No informado': ('No informado', None),
    }

    return conversiones.get(
        estado,
        ('No informado', None)
    )


def aplicar_estado_operacional(
    vehiculo,
    estado,
    subestado=None
):

    vehiculo.estado_operacional = estado

    if estado == 'No operativo':
        vehiculo.subestado_no_operativo = subestado
    else:
        vehiculo.subestado_no_operativo = None
 
INTERVALO_MANTENCION_KM = 10000

def aplicar_estado_posterior_mantencion(
    vehiculo,
    estado_posterior
):
    if estado_posterior == 'NO_CAMBIAR':

        return

    if estado_posterior == 'OPERATIVO':

        vehiculo.estado_operacional = 'Operativo'
        vehiculo.subestado_no_operativo = None

    elif estado_posterior == 'NO_INFORMADO':

        vehiculo.estado_operacional = 'No informado'
        vehiculo.subestado_no_operativo = None

    elif estado_posterior == 'NO_OPERATIVO_MANTENCION':

        vehiculo.estado_operacional = 'No operativo'
        vehiculo.subestado_no_operativo = 'Mantencion'

    elif estado_posterior == 'NO_OPERATIVO_REPARACION':

        vehiculo.estado_operacional = 'No operativo'
        vehiculo.subestado_no_operativo = 'Reparacion'

    elif estado_posterior == 'NO_OPERATIVO_DETENIDO':

        vehiculo.estado_operacional = 'No operativo'
        vehiculo.subestado_no_operativo = 'Detenido'
        
def actualizar_ciclo_mantencion_km(
    vehiculo,
    kilometraje_cierre
):
    vehiculo.kilometraje_ultima_mantencion = kilometraje_cierre

    vehiculo.kilometraje_proxima_mantencion = (
        kilometraje_cierre + INTERVALO_MANTENCION_KM
    )

    if vehiculo.kilometraje_actual is not None:

        vehiculo.kilometraje_restante = (
            vehiculo.kilometraje_proxima_mantencion
            - vehiculo.kilometraje_actual
        )

        if vehiculo.kilometraje_restante <= 0:

            vehiculo.alerta_mantencion = 'Vencida'

        elif vehiculo.kilometraje_restante <= 1500:

            vehiculo.alerta_mantencion = 'Próxima'

        else:

            vehiculo.alerta_mantencion = 'Al día'

    else:

        vehiculo.kilometraje_restante = None
        vehiculo.alerta_mantencion = 'Sin información'

def obtener_kilometraje_programado_mantencion(vehiculo):
    
    if vehiculo.kilometraje_ultima_mantencion is not None:

        return vehiculo.kilometraje_ultima_mantencion + INTERVALO_MANTENCION_KM

    if vehiculo.kilometraje_proxima_mantencion is not None:

        return vehiculo.kilometraje_proxima_mantencion

    return None
    
def calcular_alerta_kilometraje_mantencion(vehiculo):

    kilometraje_actual = vehiculo.kilometraje_actual

    kilometraje_programado = obtener_kilometraje_programado_mantencion(
        vehiculo
    )

    problemas = []

    if kilometraje_actual is None:

        problemas.append(
            'Sin kilometraje actual'
        )

    if (
        vehiculo.kilometraje_ultima_mantencion is None
        and vehiculo.kilometraje_proxima_mantencion is None
    ):

        problemas.append(
            'Sin última ni próxima mantención'
        )

    elif vehiculo.kilometraje_ultima_mantencion is None:

        problemas.append(
            'Sin última mantención'
        )

    elif vehiculo.kilometraje_proxima_mantencion is None:

        problemas.append(
            'Sin próxima mantención registrada'
        )

    if kilometraje_actual is None or kilometraje_programado is None:

        return {
            'vehiculo': vehiculo,
            'kilometraje_actual': kilometraje_actual,
            'kilometraje_programado': kilometraje_programado,
            'kilometros_restantes': None,
            'estado_alerta': 'SIN_DATOS',
            'problemas': problemas if problemas else ['Sin datos suficientes'],
        }

    kilometros_restantes = kilometraje_programado - kilometraje_actual

    if kilometros_restantes <= 0:

        estado_alerta = 'VENCIDA'

    elif kilometros_restantes <= 1500:

        estado_alerta = 'PROXIMA'

    else:

        estado_alerta = 'AL_DIA'

    return {
        'vehiculo': vehiculo,
        'kilometraje_actual': kilometraje_actual,
        'kilometraje_programado': kilometraje_programado,
        'kilometros_restantes': kilometros_restantes,
        'estado_alerta': estado_alerta,
        'problemas': problemas,
    }

def obtener_mantencion_en_curso_vehiculo(vehiculo):
    return vehiculo.mantenciones.filter(
        estado='EN_CURSO'
    ).order_by(
        'fecha_ingreso',
        'id'
    ).first()

def calcular_porcentaje(valor, total):
    if not total:
        return 0

    return round(
        (valor / total) * 100,
        1
    )


def obtener_resumen_mantenciones_km():

    vehiculos = Vehiculo.objects.select_related(
        'centro_costo'
    ).exclude(
        estado_administrativo='Dado de Baja'
    ).order_by(
        'patente'
    )

    alertas_km_todas = []

    for vehiculo in vehiculos:

        alerta = calcular_alerta_kilometraje_mantencion(
            vehiculo
        )

        alertas_km_todas.append(
            alerta
        )

    alertas_km_vencidas = [
        alerta for alerta in alertas_km_todas
        if alerta['estado_alerta'] == 'VENCIDA'
    ]

    alertas_km_proximas = [
        alerta for alerta in alertas_km_todas
        if alerta['estado_alerta'] == 'PROXIMA'
    ]

    alertas_km_al_dia = [
        alerta for alerta in alertas_km_todas
        if alerta['estado_alerta'] == 'AL_DIA'
    ]

    alertas_km_sin_datos = [
        alerta for alerta in alertas_km_todas
        if alerta['estado_alerta'] == 'SIN_DATOS'
    ]

    total_vehiculos_controlados = vehiculos.count()

    total_km_al_dia = len(
        alertas_km_al_dia
    )

    total_km_proximas = len(
        alertas_km_proximas
    )

    total_km_vencidas = len(
        alertas_km_vencidas
    )

    total_km_sin_datos = len(
        alertas_km_sin_datos
    )

    def calcular_porcentaje(cantidad):

        if not total_vehiculos_controlados:

            return 0

        return round(
            (cantidad / total_vehiculos_controlados) * 100,
            1
        )

    return {
        'vehiculos_controlados': vehiculos,

        'alertas_km_todas': alertas_km_todas,
        'alertas_km_vencidas': alertas_km_vencidas,
        'alertas_km_proximas': alertas_km_proximas,
        'alertas_km_al_dia': alertas_km_al_dia,
        'alertas_km_sin_datos': alertas_km_sin_datos,

        'total_vehiculos_controlados': total_vehiculos_controlados,

        'total_km_al_dia': total_km_al_dia,
        'total_km_proximas': total_km_proximas,
        'total_km_vencidas': total_km_vencidas,
        'total_km_sin_datos': total_km_sin_datos,

        'porcentaje_km_al_dia': calcular_porcentaje(
            total_km_al_dia
        ),

        'porcentaje_km_proximas': calcular_porcentaje(
            total_km_proximas
        ),

        'porcentaje_km_vencidas': calcular_porcentaje(
            total_km_vencidas
        ),

        'porcentaje_km_sin_datos': calcular_porcentaje(
            total_km_sin_datos
        ),
    }

@login_required
@editor_required
def iniciar_mantencion_vehiculo(request, id):

    mantencion = MantencionVehiculo.objects.select_related(
        'vehiculo',
        'vehiculo__centro_costo'
    ).get(
        id=id
    )

    vehiculo = mantencion.vehiculo

    if mantencion.estado == 'EN_CURSO':

        messages.info(
            request,
            'Esta mantención ya se encuentra en curso.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if mantencion.estado == 'CERRADA':

        messages.error(
            request,
            'No puedes iniciar una mantención que ya está cerrada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if mantencion.estado == 'CANCELADA':

        messages.error(
            request,
            'No puedes iniciar una mantención cancelada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    mantencion_en_curso = obtener_mantencion_en_curso_vehiculo(
        vehiculo
    )

    if mantencion_en_curso:

        messages.error(
            request,
            (
                f'No se puede iniciar esta mantención porque el vehículo '
                f'{vehiculo.patente} ya tiene una mantención en curso. '
                f'Debes cerrar la mantención activa antes de iniciar otra.'
            )
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    estado_mantencion_anterior = (
        f"Estado mantención: {mantencion.estado} | "
        f"Fecha ingreso: {mantencion.fecha_ingreso} | "
        f"Kilometraje ingreso: {mantencion.kilometraje_ingreso}"
    )

    estado_vehiculo_anterior = (
        f"Estado operacional: {vehiculo.estado_operacional} | "
        f"Subestado: {vehiculo.subestado_no_operativo}"
    )

    mantencion.estado = 'EN_CURSO'

    if not mantencion.fecha_ingreso:

        mantencion.fecha_ingreso = timezone.now().date()

    if (
        not mantencion.kilometraje_ingreso
        and vehiculo.kilometraje_actual
    ):

        mantencion.kilometraje_ingreso = vehiculo.kilometraje_actual

    mantencion.save()

    vehiculo.estado_operacional = 'No operativo'

    if mantencion.tipo_mantencion in [
        'PROGRAMADA',
        'KILOMETRAJE',
    ]:

        vehiculo.subestado_no_operativo = 'Mantencion'

    elif mantencion.tipo_mantencion == 'SINIESTRO':

        vehiculo.subestado_no_operativo = 'Reparacion'

    vehiculo.save()

    estado_mantencion_nuevo = (
        f"Estado mantención: {mantencion.estado} | "
        f"Fecha ingreso: {mantencion.fecha_ingreso} | "
        f"Kilometraje ingreso: {mantencion.kilometraje_ingreso}"
    )

    estado_vehiculo_nuevo = (
        f"Estado operacional: {vehiculo.estado_operacional} | "
        f"Subestado: {vehiculo.subestado_no_operativo}"
    )

    registrar_bitacora(
        request=request,
        accion='CAMBIO_ESTADO',
        modulo='Mantenciones',
        modelo_afectado='MantencionVehiculo',
        objeto_id=mantencion.id,
        objeto_repr=(
            f'{vehiculo.patente} - '
            f'{mantencion.get_tipo_mantencion_display()}'
        ),
        descripcion=(
            f'Mantención de vehículo {vehiculo.patente} iniciada.'
        ),
        valor_anterior=estado_mantencion_anterior,
        valor_nuevo=estado_mantencion_nuevo
    )

    registrar_bitacora(
        request=request,
        accion='CAMBIO_ESTADO',
        modulo='Vehículos',
        modelo_afectado='Vehiculo',
        objeto_id=vehiculo.id,
        objeto_repr=vehiculo.patente,
        descripcion=(
            f'Vehículo {vehiculo.patente} actualizado automáticamente '
            f'por inicio de mantención.'
        ),
        valor_anterior=estado_vehiculo_anterior,
        valor_nuevo=estado_vehiculo_nuevo
    )

    if mantencion.tipo_mantencion == 'SINIESTRO':

        messages.success(
            request,
            (
                f'Mantención iniciada correctamente. '
                f'El vehículo {vehiculo.patente} fue actualizado a '
                f'No operativo / Reparación.'
            )
        )

    else:

        messages.success(
            request,
            (
                f'Mantención iniciada correctamente. '
                f'El vehículo {vehiculo.patente} fue actualizado a '
                f'No operativo / Mantención.'
            )
        )

    return redirect(
        'detalle_vehiculo',
        id=vehiculo.id
    )

@login_required
@editor_required
def cancelar_mantencion_vehiculo(request, id):

    mantencion = get_object_or_404(
        MantencionVehiculo.objects.select_related(
            'vehiculo',
            'vehiculo__centro_costo',
            'usuario_registro',
            'usuario_cierre'
        ),
        id=id
    )

    vehiculo = mantencion.vehiculo

    if mantencion.estado == 'CERRADA':

        messages.error(
            request,
            'No puedes cancelar una mantención que ya está cerrada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if mantencion.estado == 'CANCELADA':

        messages.info(
            request,
            'Esta mantención ya se encuentra cancelada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if mantencion.estado == 'EN_CURSO':

        messages.error(
            request,
            (
                'No puedes cancelar una mantención que ya está en curso. '
                'Si la mantención se ejecutó, debes cerrarla. '
                'Si fue iniciada por error, revisa el estado del vehículo antes de continuar.'
            )
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if request.method == 'POST':

        form = CancelarMantencionVehiculoForm(
            request.POST
        )

        if form.is_valid():

            estado_anterior = (
                f"Estado: {mantencion.estado} | "
                f"Fecha programada: {mantencion.fecha_programada} | "
                f"Fecha cierre/cancelación: {mantencion.fecha_cierre} | "
                f"Observación cierre/cancelación: {mantencion.observacion_cierre}"
            )

            observacion_cancelacion = form.cleaned_data[
                'observacion_cancelacion'
            ]

            mantencion.estado = 'CANCELADA'
            mantencion.fecha_cierre = timezone.now().date()
            mantencion.observacion_cierre = observacion_cancelacion
            mantencion.usuario_cierre = request.user

            mantencion.save()

            estado_nuevo = (
                f"Estado: {mantencion.estado} | "
                f"Fecha programada: {mantencion.fecha_programada} | "
                f"Fecha cierre/cancelación: {mantencion.fecha_cierre} | "
                f"Observación cierre/cancelación: {mantencion.observacion_cierre}"
            )

            registrar_bitacora(
                request=request,
                accion='EDITAR',
                modulo='Mantenciones',
                modelo_afectado='MantencionVehiculo',
                objeto_id=mantencion.id,
                objeto_repr=(
                    f'{vehiculo.patente} - '
                    f'{mantencion.get_tipo_mantencion_display()}'
                ),
                descripcion=(
                    f'Mantención pendiente de vehículo {vehiculo.patente} cancelada.'
                ),
                valor_anterior=estado_anterior,
                valor_nuevo=estado_nuevo
            )

            messages.success(
                request,
                f'Mantención de {vehiculo.patente} cancelada correctamente.'
            )

            return redirect(
                'detalle_vehiculo',
                id=vehiculo.id
            )

    else:

        form = CancelarMantencionVehiculoForm()

    return render(
        request,
        'flota/cancelar_mantencion.html',
        {
            'form': form,
            'mantencion': mantencion,
            'vehiculo': vehiculo,
        }
    )

@login_required
@editor_required
def reprogramar_mantencion_vehiculo(request, id):

    mantencion = get_object_or_404(
        MantencionVehiculo.objects.select_related(
            'vehiculo',
            'vehiculo__centro_costo',
            'usuario_registro'
        ),
        id=id
    )

    vehiculo = mantencion.vehiculo

    if mantencion.estado == 'CERRADA':

        messages.error(
            request,
            'No puedes reprogramar una mantención cerrada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if mantencion.estado == 'CANCELADA':

        messages.error(
            request,
            'No puedes reprogramar una mantención cancelada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if mantencion.estado == 'EN_CURSO':

        messages.error(
            request,
            'No puedes reprogramar una mantención que ya está en curso.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if request.method == 'POST':

        form = ReprogramarMantencionVehiculoForm(
            request.POST,
            mantencion=mantencion
        )

        if form.is_valid():

            estado_anterior = (
                f"Fecha programada: {mantencion.fecha_programada} | "
                f"Motivo: {mantencion.motivo} | "
                f"Observación: {mantencion.observacion}"
            )

            nueva_fecha_programada = form.cleaned_data[
                'nueva_fecha_programada'
            ]

            motivo = form.cleaned_data[
                'motivo'
            ]

            observacion = form.cleaned_data[
                'observacion'
            ]

            mantencion.fecha_programada = nueva_fecha_programada
            mantencion.motivo = motivo
            mantencion.observacion = observacion

            mantencion.save()

            estado_nuevo = (
                f"Fecha programada: {mantencion.fecha_programada} | "
                f"Motivo: {mantencion.motivo} | "
                f"Observación: {mantencion.observacion}"
            )

            registrar_bitacora(
                request=request,
                accion='EDITAR',
                modulo='Mantenciones',
                modelo_afectado='MantencionVehiculo',
                objeto_id=mantencion.id,
                objeto_repr=(
                    f'{vehiculo.patente} - '
                    f'{mantencion.get_tipo_mantencion_display()}'
                ),
                descripcion=(
                    f'Mantención de vehículo {vehiculo.patente} reprogramada.'
                ),
                valor_anterior=estado_anterior,
                valor_nuevo=estado_nuevo
            )

            messages.success(
                request,
                f'Mantención de {vehiculo.patente} reprogramada correctamente.'
            )

            return redirect(
                'detalle_vehiculo',
                id=vehiculo.id
            )

    else:

        form = ReprogramarMantencionVehiculoForm(
            initial={
                'nueva_fecha_programada':
                mantencion.fecha_programada.isoformat()
                if mantencion.fecha_programada
                else timezone.now().date().isoformat(),

                'motivo':
                mantencion.motivo,

                'observacion':
                mantencion.observacion,
            },
            mantencion=mantencion
        )

    return render(
        request,
        'flota/reprogramar_mantencion.html',
        {
            'form': form,
            'mantencion': mantencion,
            'vehiculo': vehiculo,
        }
    )

@login_required
@editor_required
def cerrar_mantencion_vehiculo(request, id):

    mantencion = MantencionVehiculo.objects.select_related(
        'vehiculo',
        'vehiculo__centro_costo'
    ).get(
        id=id
    )

    vehiculo = mantencion.vehiculo

    if mantencion.estado == 'CERRADA':

        messages.info(
            request,
            'Esta mantención ya se encuentra cerrada.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    if request.method == 'POST':

        form = CerrarMantencionVehiculoForm(
            request.POST,
            mantencion=mantencion
        )

        if form.is_valid():

            estado_mantencion_anterior = (
                f"Estado mantención: {mantencion.estado} | "
                f"Fecha cierre: {mantencion.fecha_cierre} | "
                f"Kilometraje cierre: {mantencion.kilometraje_cierre} | "
                f"Observación cierre: {mantencion.observacion_cierre}"
            )

            estado_vehiculo_anterior = (
                f"Estado operacional: {vehiculo.estado_operacional} | "
                f"Subestado: {vehiculo.subestado_no_operativo} | "
                f"Última mantención: {vehiculo.kilometraje_ultima_mantencion} | "
                f"Próxima mantención: {vehiculo.kilometraje_proxima_mantencion} | "
                f"Km restantes: {vehiculo.kilometraje_restante} | "
                f"Alerta mantención: {vehiculo.alerta_mantencion}"
            )

            fecha_cierre = form.cleaned_data[
                'fecha_cierre'
            ]

            kilometraje_cierre = form.cleaned_data[
                'kilometraje_cierre'
            ]

            observacion_cierre = form.cleaned_data[
                'observacion_cierre'
            ]

            estado_posterior = form.cleaned_data[
                'estado_posterior_vehiculo'
            ]

            mantencion.estado = 'CERRADA'
            mantencion.fecha_cierre = fecha_cierre
            mantencion.kilometraje_cierre = kilometraje_cierre
            mantencion.observacion_cierre = observacion_cierre
            mantencion.usuario_cierre = request.user

            mantencion.save()

            if mantencion.tipo_mantencion == 'KILOMETRAJE':

                actualizar_ciclo_mantencion_km(
                    vehiculo,
                    kilometraje_cierre
                )

            aplicar_estado_posterior_mantencion(
                vehiculo,
                estado_posterior
            )

            vehiculo.save()

            estado_mantencion_nuevo = (
                f"Estado mantención: {mantencion.estado} | "
                f"Fecha cierre: {mantencion.fecha_cierre} | "
                f"Kilometraje cierre: {mantencion.kilometraje_cierre} | "
                f"Observación cierre: {mantencion.observacion_cierre}"
            )

            estado_vehiculo_nuevo = (
                f"Estado operacional: {vehiculo.estado_operacional} | "
                f"Subestado: {vehiculo.subestado_no_operativo} | "
                f"Última mantención: {vehiculo.kilometraje_ultima_mantencion} | "
                f"Próxima mantención: {vehiculo.kilometraje_proxima_mantencion} | "
                f"Km restantes: {vehiculo.kilometraje_restante} | "
                f"Alerta mantención: {vehiculo.alerta_mantencion}"
            )

            registrar_bitacora(
                request=request,
                accion='EDITAR',
                modulo='Mantenciones',
                modelo_afectado='MantencionVehiculo',
                objeto_id=mantencion.id,
                objeto_repr=(
                    f'{vehiculo.patente} - '
                    f'{mantencion.get_tipo_mantencion_display()}'
                ),
                descripcion=(
                    f'Cierre de mantención para vehículo '
                    f'{vehiculo.patente}.'
                ),
                valor_anterior=estado_mantencion_anterior,
                valor_nuevo=estado_mantencion_nuevo
            )

            registrar_bitacora(
                request=request,
                accion='CAMBIO_ESTADO',
                modulo='Vehículos',
                modelo_afectado='Vehiculo',
                objeto_id=vehiculo.id,
                objeto_repr=vehiculo.patente,
                descripcion=(
                    f'Vehículo {vehiculo.patente} actualizado por cierre '
                    f'de mantención.'
                ),
                valor_anterior=estado_vehiculo_anterior,
                valor_nuevo=estado_vehiculo_nuevo
            )

            messages.success(
                request,
                f'Mantención de {vehiculo.patente} cerrada correctamente.'
            )

            return redirect(
                'detalle_vehiculo',
                id=vehiculo.id
            )

    else:

        form = CerrarMantencionVehiculoForm(
            initial={
                'fecha_cierre': timezone.now().date().isoformat(),
                'kilometraje_cierre': vehiculo.kilometraje_actual,
                'estado_posterior_vehiculo': 'OPERATIVO',
            },
            mantencion=mantencion
        )

    return render(
        request,
        'flota/cerrar_mantencion.html',
        {
            'form': form,
            'mantencion': mantencion,
            'vehiculo': vehiculo,
        }
    )


def obtener_metricas_estado(queryset):

    total_vehiculos = queryset.count()

    operativos = queryset.filter(
        estado_operacional='Operativo'
    ).count()

    no_operativos = queryset.filter(
        estado_operacional='No operativo'
    ).count()

    no_informados = queryset.filter(
        estado_operacional='No informado'
    ).count()

    en_mantencion = queryset.filter(
        estado_operacional='No operativo',
        subestado_no_operativo='Mantencion'
    ).count()

    en_reparacion = queryset.filter(
        estado_operacional='No operativo',
        subestado_no_operativo='Reparacion'
    ).count()

    detenidos = queryset.filter(
        estado_operacional='No operativo',
        subestado_no_operativo='Detenido'
    ).count()
    
    porcentaje_mantencion_no_operativo = round(
    (en_mantencion / no_operativos) * 100,
    1
    ) if no_operativos else 0

    porcentaje_reparacion_no_operativo = round(
        (en_reparacion / no_operativos) * 100,
        1
    ) if no_operativos else 0

    porcentaje_detenidos_no_operativo = round(
        (detenidos / no_operativos) * 100,
        1
    ) if no_operativos else 0

    porcentaje_operativos = round(
        (operativos / total_vehiculos) * 100,
        1
    ) if total_vehiculos else 0

    porcentaje_no_operativos = round(
        (no_operativos / total_vehiculos) * 100,
        1
    ) if total_vehiculos else 0

    porcentaje_no_informado = round(
        (no_informados / total_vehiculos) * 100,
        1
    ) if total_vehiculos else 0

    return {
        'total_vehiculos': total_vehiculos,
        'operativos': operativos,
        'no_operativos': no_operativos,
        'no_informados': no_informados,
        'en_mantencion': en_mantencion,
        'en_reparacion': en_reparacion,
        'detenidos': detenidos,
        'porcentaje_mantencion_no_operativo': porcentaje_mantencion_no_operativo,
        'porcentaje_reparacion_no_operativo': porcentaje_reparacion_no_operativo,
        'porcentaje_detenidos_no_operativo': porcentaje_detenidos_no_operativo,
        'porcentaje_operativos': porcentaje_operativos,
        'porcentaje_no_operativos': porcentaje_no_operativos,
        'porcentaje_no_informado': porcentaje_no_informado,
       

        # Compatibilidad temporal con templates antiguos.
        'fuera_servicio': no_operativos,
        'porcentaje_fuera': porcentaje_no_operativos,
        'porcentaje_mantencion': round(
            (en_mantencion / total_vehiculos) * 100,
            1
        ) if total_vehiculos else 0,
    }
@login_required
def inicio(request):

    vehiculos_base = Vehiculo.objects.exclude(
        estado_administrativo='Dado de Baja'
    )

    metricas = obtener_metricas_estado(
        vehiculos_base
    )

    metricas_mantenciones = obtener_resumen_mantenciones_km()

    metricas.update(
        metricas_mantenciones
    )

    return render(
        request,
        'flota/inicio.html',
        metricas
    )

@login_required
def dashboard(request):

    centro_costo_id = request.GET.get('centro_costo')
    marca = request.GET.get('marca')
    modelo = request.GET.get('modelo')
    anio = request.GET.get('anio')
    estado = request.GET.get('estado')
    subestado = request.GET.get('subestado')
    patente = request.GET.get('patente','').strip().upper()
    orden = request.GET.get('orden', 'patente')

    mostrar_bajas = request.GET.get(
        'mostrar_bajas'
    )

    vehiculos_base = Vehiculo.objects.select_related(
        'centro_costo'
    ).all()

    if not mostrar_bajas:

        vehiculos_base = vehiculos_base.exclude(
            estado_administrativo='Dado de Baja'
        )

    vehiculos_filtrados = vehiculos_base

    if centro_costo_id:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            centro_costo_id=centro_costo_id
        )

    if marca:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            marca=marca
        )

    if modelo:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            modelo=modelo
        )

    if anio:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            anio=anio
        )

    if estado:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            estado_operacional=estado
        )

    if subestado:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            subestado_no_operativo=subestado
        )

    if patente:
        vehiculos_filtrados = vehiculos_filtrados.filter(
            patente__icontains=patente
        )

    opciones_base = vehiculos_base

    if centro_costo_id:
        opciones_base = opciones_base.filter(
            centro_costo_id=centro_costo_id
        )

    marcas = opciones_base.values_list(
        'marca',
        flat=True
    ).exclude(
        marca__isnull=True
    ).distinct().order_by('marca')

    if marca:
        opciones_base = opciones_base.filter(
            marca=marca
        )

    modelos = opciones_base.values_list(
        'modelo',
        flat=True
    ).exclude(
        modelo__isnull=True
    ).distinct().order_by('modelo')

    if modelo:
        opciones_base = opciones_base.filter(
            modelo=modelo
        )

    anios = opciones_base.values_list(
        'anio',
        flat=True
    ).exclude(
        anio__isnull=True
    ).distinct().order_by('anio')

    estados = Vehiculo.ESTADOS_OPERACIONALES

    metricas = obtener_metricas_estado(
        vehiculos_filtrados
    )

    ordenes_permitidos = [
        'patente',
        'kilometraje_actual',
        '-kilometraje_actual',
    ]

    if orden in ordenes_permitidos:

        vehiculos_filtrados = vehiculos_filtrados.order_by(
            orden
        )

    else:

        vehiculos_filtrados = vehiculos_filtrados.order_by(
            'patente'
        )

    paginator = Paginator(
        vehiculos_filtrados,
        20
    )

    page_number = request.GET.get(
        'page'
    )

    vehiculos = paginator.get_page(
        page_number
    )

    query_params = request.GET.copy()

    if 'page' in query_params:
        query_params.pop('page')

    if 'orden' in query_params:
        query_params.pop('orden')

    query_string = query_params.urlencode()

    contexto = {
        'vehiculos': vehiculos,

        'centros_costo':
        CentroCosto.objects.order_by(
            'codigo'
        ),

        'marcas': marcas,
        'modelos': modelos,
        'anios': anios,
        'estados': estados,

        'subestados_no_operativo':
        Vehiculo.SUBESTADOS_NO_OPERATIVO,

        'centro_costo_id':
        centro_costo_id,

        'marca_seleccionada':
        marca,

        'modelo_seleccionado':
        modelo,

        'anio_seleccionado':
        anio,

        'estado_seleccionado':
        estado,

        'subestado_seleccionado':
        subestado,

        'patente_buscada':
        patente,

        'mostrar_bajas':
        mostrar_bajas,

        'orden_actual':
        orden,

        'query_string':
        query_string,
    }

    contexto.update(
        metricas
    )

    return render(
        request,
        'flota/dashboard.html',
        contexto
    )


def exportar_vehiculos_excel(request):
    vehiculos = Vehiculo.objects.select_related('centro_costo').all()

    centro_costo_id = request.GET.get('centro_costo')
    marca = request.GET.get('marca')
    modelo = request.GET.get('modelo')
    anio = request.GET.get('anio')
    estado = request.GET.get('estado')
    subestado = request.GET.get('subestado')
    patente = request.GET.get('patente','').strip().upper()

    if centro_costo_id:
        vehiculos = vehiculos.filter(
            centro_costo_id=centro_costo_id
        )

    if marca:

        vehiculos = vehiculos.annotate(
            marca_limpia=Upper(
                Trim(
                    'marca'
                )
            )
        ).filter(
            marca_limpia=marca
        )

    if modelo:
        vehiculos = vehiculos.filter(
            modelo=modelo
        )

    if anio:
        vehiculos = vehiculos.filter(
            anio=anio
        )

    if estado:
        vehiculos = vehiculos.filter(
            estado_operacional=estado
        )

    if subestado:
        vehiculos = vehiculos.filter(
            subestado_no_operativo=subestado
        )

    if patente:
        vehiculos = vehiculos.filter(
            patente__icontains=patente
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Vehículos"

    ws.append([
        "Patente",
        "Año",
        "Marca",
        "Modelo",
        "Tipo Vehículo",
        "Centro de costo",
        "Estado Operacional",
        "Subestado No Operativo",
        "Estado Administrativo",
        "Estado Mantención",
        "Kilometraje actual",
        "Kilometraje próxima mantención",
        "Kilometraje restante",
        "Alerta mantención",
        "Empresa arrendadora",
        "Faena",
        "Gerencia",
        "Fecha baja",
    ])

    for vehiculo in vehiculos:
        ws.append([
            vehiculo.patente,
            vehiculo.anio,
            vehiculo.marca,
            vehiculo.modelo,
            vehiculo.tipo_vehiculo,
            str(vehiculo.centro_costo) if vehiculo.centro_costo else "",
            vehiculo.estado_operacional,
            vehiculo.get_subestado_no_operativo_display()
            if vehiculo.subestado_no_operativo
            else "",
            vehiculo.estado_administrativo,
            vehiculo.estado_mantencion,
            vehiculo.kilometraje_actual,
            vehiculo.kilometraje_proxima_mantencion,
            vehiculo.kilometraje_restante,
            vehiculo.alerta_mantencion,
            vehiculo.empresa_arrendadora,
            vehiculo.faena,
            vehiculo.gerencia,
            vehiculo.fecha_baja,
        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="vehiculos_exportados.xlsx"'
    )

    wb.save(response)

    return response

@login_required
@master_required
def importar_datos(request):
    mensaje = None
    detalle_importacion = []
    total_alertas = 0
    archivo_temporal = None

    if request.method == 'POST':
        archivo = request.FILES.get('archivo_excel')

        if archivo:
            archivo_temporal = default_storage.save(
                f"importaciones/{archivo.name}",
                ContentFile(archivo.read())
            )

            archivo.seek(0)

            wb = load_workbook(archivo)
            hoja = wb.active

            encabezados_excel = obtener_encabezados(hoja)

            columnas_obligatorias = [
                'Patente',
                'Centro De Costo',
                'Kilometraje De La Camioneta',
                'Fecha En Que Se Dio De Baja',
            ]

            faltantes = []

            for columna in columnas_obligatorias:
                if columna not in encabezados_excel:
                    faltantes.append(columna)

            if faltantes:
                mensaje = (
                    "Faltan columnas obligatorias: "
                    + ", ".join(faltantes)
                )

            else:
                errores = []
                patentes_vistas = set()

                for fila_numero, fila in enumerate(
                    hoja.iter_rows(min_row=3, values_only=True),
                    start=3
                ):
                    datos = obtener_datos_fila(encabezados_excel, fila)

                    patente = limpiar_texto(
                        datos.get('patente')
                    )

                    if patente:
                        patente_limpia = patente.upper()

                        if patente_limpia in patentes_vistas:
                            errores.append(
                                f"Fila {fila_numero}: patente duplicada "
                                f"en Excel ({patente_limpia})"
                            )
                        else:
                            patentes_vistas.add(patente_limpia)

                    campos_obligatorios = [
                        'patente',
                        'centro_costo',
                    ]

                    for campo in campos_obligatorios:
                        valor = datos.get(campo)

                        if valor is None or str(valor).strip() == "":
                            errores.append(
                                f"Fila {fila_numero}: {campo} vacío"
                            )

                if errores:
                    mensaje = (
                        f"<strong>Errores encontrados: "
                        f"{len(errores)}</strong><br><br>"
                        + "<br>".join(errores)
                    )

                else:
                    nuevos = 0
                    actualizados = 0
                    sin_cambios = 0
                    total_alertas = 0
                    detalle_importacion = []

                    for fila in hoja.iter_rows(
                        min_row=3,
                        values_only=True
                    ):
                        datos = obtener_datos_fila(encabezados_excel, fila)

                        patente = limpiar_texto(
                            datos.get('patente')
                        ).upper()

                        vehiculo = Vehiculo.objects.filter(
                            patente=patente
                        ).first()

                        centro_costo_codigo = limpiar_texto(
                            datos.get('centro_costo')
                        )

                        fecha_baja = limpiar_fecha(
                            datos.get('fecha_baja')
                        )

                        estado_administrativo = obtener_estado_administrativo(
                            fecha_baja
                        )

                        if not vehiculo:
                            nuevos += 1

                            detalle_importacion.append({
                                'centro_costo': centro_costo_codigo,
                                'patente': patente,
                                'accion': 'Nuevo',
                                'tiene_alerta': False,
                                'cambios': 'Registro nuevo'
                            })

                        else:
                            cambios_detectados = []
                            alertas = []

                            campos_comparar = [
                                (
                                    'anio',
                                    'Año',
                                    limpiar_int(datos.get('anio'))
                                ),
                                (
                                    'marca',
                                    'Marca',
                                    limpiar_texto(datos.get('marca'))
                                ),
                                (
                                    'modelo',
                                    'Modelo',
                                    limpiar_texto(datos.get('modelo'))
                                ),
                                (
                                    'estado_administrativo',
                                    'Estado Administrativo',
                                    estado_administrativo
                                ),
                                (
                                    'estado_mantencion',
                                    'Estado Mantención',
                                    limpiar_texto(
                                        datos.get('estado_mantencion')
                                    )
                                ),
                                (
                                    'kilometraje_actual',
                                    'Kilometraje',
                                    limpiar_int(
                                        datos.get('kilometraje_actual')
                                    )
                                ),
                                (
                                    'faena',
                                    'Faena',
                                    limpiar_texto(datos.get('faena'))
                                ),
                                (
                                    'gerencia',
                                    'Gerencia',
                                    limpiar_texto(datos.get('gerencia'))
                                ),
                            ]

                            for campo, etiqueta, nuevo_valor in campos_comparar:
                                valor_actual = getattr(vehiculo, campo)

                                if str(valor_actual) != str(nuevo_valor):
                                    cambios_detectados.append(
                                        f"{etiqueta}: "
                                        f"{valor_actual} → {nuevo_valor}"
                                    )

                            nuevo_km = limpiar_int(
                                datos.get('kilometraje_actual')
                            )

                            km_actual = vehiculo.kilometraje_actual

                            if (
                                nuevo_km is not None
                                and km_actual is not None
                            ):
                                if nuevo_km < km_actual:
                                    alertas.append(
                                        "⚠ Kilometraje disminuye"
                                    )

                                if abs(nuevo_km - km_actual) > 50000:
                                    alertas.append(
                                        "⚠ Cambio mayor a 50.000 km"
                                    )

                            if cambios_detectados:
                                actualizados += 1

                                if alertas:
                                    total_alertas += len(alertas)

                                cambios_html = "<br>".join(
                                    cambios_detectados
                                )

                                if alertas:
                                    cambios_html += (
                                        "<br><br>"
                                        + "<br>".join(alertas)
                                    )

                                detalle_importacion.append({
                                    'centro_costo': centro_costo_codigo,
                                    'patente': patente,
                                    'accion': 'Actualizar',
                                    'tiene_alerta': len(alertas) > 0,
                                    'cambios': cambios_html
                                })

                            else:
                                sin_cambios += 1

                                detalle_importacion.append({
                                    'centro_costo': centro_costo_codigo,
                                    'patente': patente,
                                    'accion': 'Sin cambios',
                                    'tiene_alerta': False,
                                    'cambios': '-'
                                })

                    mensaje = (
                        "<strong>Validación completada</strong><br><br>"
                        f"Nuevos: {nuevos}<br>"
                        f"Actualizados: {actualizados}<br>"
                        f"Sin cambios: {sin_cambios}<br>"
                        f"Alertas detectadas: {total_alertas}"
                    )

        else:
            mensaje = "No se seleccionó archivo"

    return render(
        request,
        'flota/importar.html',
        {
            'mensaje': mensaje,
            'detalle_importacion': detalle_importacion,
            'total_alertas': total_alertas,
            'archivo_temporal': archivo_temporal
        }
    )

@login_required
@master_required
def confirmar_importacion(request):
    if request.method != 'POST':
        return HttpResponse("Método no permitido")

    archivo_temporal = request.POST.get('archivo_temporal')

    if not archivo_temporal:
        return HttpResponse("No se recibió archivo temporal")

    ruta_archivo = default_storage.path(archivo_temporal)

    wb = load_workbook(ruta_archivo)
    hoja = wb.active

    encabezados_excel = obtener_encabezados(hoja)

    nuevos = 0
    actualizados = 0
    sin_cambios = 0

    for fila in hoja.iter_rows(
        min_row=3,
        values_only=True
    ):
        datos = obtener_datos_fila(encabezados_excel, fila)

        patente = limpiar_texto(
            datos.get('patente')
        ).upper()

        centro_costo = obtener_o_crear_centro_costo(
            datos.get('centro_costo')
        )

        vehiculo = Vehiculo.objects.filter(
            patente=patente
        ).first()

        fecha_baja = limpiar_fecha(
            datos.get('fecha_baja')
        )

        estado_administrativo = obtener_estado_administrativo(
            fecha_baja
        )

        valores = {
            'anio': limpiar_int(datos.get('anio')),
            'combustible': limpiar_texto(datos.get('combustible')),
            'traccion': limpiar_texto(datos.get('traccion')),
            'numero_motor': limpiar_texto(datos.get('numero_motor')),
            'numero_chasis_vin': limpiar_texto(
                datos.get('numero_chasis_vin')
            ),
            'tipo_vehiculo': limpiar_texto(datos.get('tipo_vehiculo')),
            'marca': limpiar_texto(datos.get('marca')),
            'modelo': limpiar_texto(datos.get('modelo')),
            'color': limpiar_texto(datos.get('color')),
            'numero_asientos': limpiar_int(datos.get('numero_asientos')),
            'tarifa_mensual': limpiar_tarifa(datos.get('tarifa_mensual')),
            'centro_costo': centro_costo,
            'tipo_estandar': limpiar_texto(datos.get('tipo_estandar')),
            'fecha_inicio_contrato': limpiar_fecha(
                datos.get('fecha_inicio_contrato')
            ),
            'fecha_termino_contrato': limpiar_fecha(
                datos.get('fecha_termino_contrato')
            ),
            'periodo_meses': limpiar_int(datos.get('periodo_meses')),
            'limite_kilometraje': limpiar_int(
                datos.get('limite_kilometraje')
            ),
            'kilometraje_ultima_mantencion': limpiar_int(
                datos.get('kilometraje_ultima_mantencion')
            ),
            'kilometraje_proxima_mantencion': limpiar_int(
                datos.get('kilometraje_proxima_mantencion')
            ),
            'kilometraje_actual': limpiar_int(
                datos.get('kilometraje_actual')
            ),
            'kilometraje_restante': limpiar_int(
                datos.get('kilometraje_restante')
            ),
            'dias_restantes': limpiar_int(datos.get('dias_restantes')),
            'alerta_mantencion': limpiar_texto(
                datos.get('alerta_mantencion')
            ),
            'empresa_arrendadora': limpiar_texto(
                datos.get('empresa_arrendadora')
            ),
            'razon_social': limpiar_texto(datos.get('razon_social')),
            'rut_empresa': limpiar_texto(datos.get('rut_empresa')),
            'sucursal': limpiar_texto(datos.get('sucursal')),
            'ciudad': limpiar_texto(datos.get('ciudad')),
            'faena': limpiar_texto(datos.get('faena')),
            'gerencia': limpiar_texto(datos.get('gerencia')),
            'fecha_baja': fecha_baja,
            'estado_administrativo': estado_administrativo,
            'estado_mantencion': limpiar_texto(
                datos.get('estado_mantencion')
            ),
        }

        if estado_administrativo == 'Dado de Baja':
            valores['estado_operacional'] = 'No operativo'
            valores['subestado_no_operativo'] = 'Detenido'
        else:
            valores['subestado_no_operativo'] = None

        if not vehiculo:
            Vehiculo.objects.create(
                patente=patente,
                **valores
            )

            nuevos += 1

        else:
            cambios = False

            for campo, nuevo_valor in valores.items():
                valor_actual = getattr(vehiculo, campo)

                if str(valor_actual) != str(nuevo_valor):
                    setattr(vehiculo, campo, nuevo_valor)
                    cambios = True

            if cambios:
                vehiculo.save()
                messages.success(
                    request,
                    f"Estado de {vehiculo.patente} actualizado correctamente."
)
                actualizados += 1
            else:
                sin_cambios += 1

    if default_storage.exists(archivo_temporal):
        default_storage.delete(archivo_temporal)

    return HttpResponse(
        f"""
        <h1>Importación completada</h1>

        <p>Nuevos: {nuevos}</p>
        <p>Actualizados: {actualizados}</p>
        <p>Sin cambios: {sin_cambios}</p>

        <br>

        <a href="/vehiculos/">Volver a vehículos</a>
        """
    )
@login_required
@editor_required
def gestionar_estados(request):

    if request.method == 'POST':

        accion = request.POST.get('accion')

        if accion and accion.startswith("individual_"):

            vehiculo_id = accion.replace(
                "individual_",
                ""
            )

            nuevo_estado = request.POST.get(
                f'estado_operacional_{vehiculo_id}'
            )

            nuevo_subestado = request.POST.get(
                f'subestado_no_operativo_{vehiculo_id}'
            )

            if nuevo_estado == 'No operativo' and not nuevo_subestado:

                messages.error(
                    request,
                    "Debe seleccionar un motivo para el estado No operativo."
                )

                return redirect(
                    'gestionar_estados'
                )

            vehiculo = Vehiculo.objects.get(
                id=vehiculo_id
            )

            estado_anterior = (
                f"Estado operacional: {vehiculo.estado_operacional} | "
                f"Subestado: {vehiculo.subestado_no_operativo}"
            )

            aplicar_estado_operacional(
                vehiculo,
                nuevo_estado,
                nuevo_subestado
            )

            estado_nuevo = (
                f"Estado operacional: {vehiculo.estado_operacional} | "
                f"Subestado: {vehiculo.subestado_no_operativo}"
            )

            vehiculo.save()

            registrar_bitacora(
                request=request,
                accion='CAMBIO_ESTADO',
                modulo='Gestión de Estados',
                modelo_afectado='Vehiculo',
                objeto_id=vehiculo.id,
                objeto_repr=vehiculo.patente,
                descripcion=f'Cambio individual de estado del vehículo {vehiculo.patente}.',
                valor_anterior=estado_anterior,
                valor_nuevo=estado_nuevo
            )

            messages.success(
                request,
                f"Estado de {vehiculo.patente} actualizado correctamente."
            )

        elif accion == "masivo":

            ids = request.POST.getlist(
                'vehiculos'
            )

            nuevo_estado = request.POST.get(
                'estado_masivo'
            )

            nuevo_subestado = request.POST.get(
                'subestado_masivo'
            )

            if ids and nuevo_estado:

                if nuevo_estado == 'No operativo' and not nuevo_subestado:

                    messages.error(
                        request,
                        "Debe seleccionar un motivo para aplicar el estado No operativo."
                    )

                    return redirect(
                        'gestionar_estados'
                    )

                vehiculos_a_cambiar = Vehiculo.objects.filter(
                    id__in=ids
                )

                cantidad = vehiculos_a_cambiar.count()

                for vehiculo in vehiculos_a_cambiar:

                    estado_anterior = (
                        f"Estado operacional: {vehiculo.estado_operacional} | "
                        f"Subestado: {vehiculo.subestado_no_operativo}"
                    )

                    aplicar_estado_operacional(
                        vehiculo,
                        nuevo_estado,
                        nuevo_subestado
                    )

                    estado_nuevo = (
                        f"Estado operacional: {vehiculo.estado_operacional} | "
                        f"Subestado: {vehiculo.subestado_no_operativo}"
                    )

                    vehiculo.save()

                    registrar_bitacora(
                        request=request,
                        accion='CAMBIO_ESTADO',
                        modulo='Gestión de Estados',
                        modelo_afectado='Vehiculo',
                        objeto_id=vehiculo.id,
                        objeto_repr=vehiculo.patente,
                        descripcion=(
                            f'Cambio masivo de estado aplicado al vehículo '
                            f'{vehiculo.patente}.'
                        ),
                        valor_anterior=estado_anterior,
                        valor_nuevo=estado_nuevo
                    )

                if cantidad > 0:

                    messages.success(
                        request,
                        f"{cantidad} vehículos actualizados correctamente."
                    )

                else:

                    messages.info(
                        request,
                        "No se seleccionaron vehículos para actualizar."
                    )

    patente = request.GET.get('patente','').strip().upper()

    centro_costo_id = request.GET.get(
        'centro_costo'
    )

    estado = request.GET.get(
        'estado'
    )

    subestado = request.GET.get(
        'subestado'
    )

    vehiculos = Vehiculo.objects.select_related(
        'centro_costo'
    ).exclude(
        estado_administrativo='Dado de Baja'
    )

    if patente:

        vehiculos = vehiculos.filter(
            patente__icontains=patente
        )

    if centro_costo_id:

        vehiculos = vehiculos.filter(
            centro_costo_id=centro_costo_id
        )

    if estado:

        vehiculos = vehiculos.filter(
            estado_operacional=estado
        )

    if subestado:

        vehiculos = vehiculos.filter(
            subestado_no_operativo=subestado
        )

    paginator = Paginator(
        vehiculos,
        25
    )

    page_number = request.GET.get(
        'page'
    )

    vehiculos_pagina = paginator.get_page(
        page_number
    )

    query_params = request.GET.copy()

    if 'page' in query_params:
        query_params.pop('page')

    query_string = query_params.urlencode()

    contexto = {
        'vehiculos': vehiculos_pagina,
        'query_string': query_string,

        'centros_costo':
        CentroCosto.objects.order_by(
            'codigo'
        ),

        'estados_operacionales':
        Vehiculo.ESTADOS_OPERACIONALES,

        'subestados_no_operativo':
        Vehiculo.SUBESTADOS_NO_OPERATIVO,

        'patente_buscada': patente,

        'centro_costo_id':
        centro_costo_id,

        'estado_seleccionado':
        estado,

        'subestado_seleccionado':
        subestado,
    }

    return render(
        request,
        'flota/estados.html',
        contexto
    )
@login_required
def detalle_vehiculo(request, id):

    vehiculo = Vehiculo.objects.select_related(
        'centro_costo'
    ).get(
        id=id
    )

    historial_bajas = vehiculo.historial_bajas.select_related(
        'usuario_baja',
        'usuario_reactivacion'
    ).all()

    historial_transferencias = vehiculo.historial_transferencias.select_related(
        'centro_costo_origen',
        'centro_costo_destino',
        'usuario_transferencia'
    ).all()

    historial_mantenciones = vehiculo.mantenciones.select_related(
        'usuario_registro',
        'usuario_cierre'
    ).all()

    contexto = {
        'vehiculo': vehiculo,
        'historial_bajas': historial_bajas,
        'historial_transferencias': historial_transferencias,
        'historial_mantenciones': historial_mantenciones,
        'puede_editar': (
            request.user.groups.filter(name='Editor').exists()
            or request.user.groups.filter(name='Master').exists()
        ),
    }

    return render(
        request,
        'flota/detalle_vehiculo.html',
        contexto
    )
    
@login_required
@editor_required
def crear_vehiculo(request):

    if request.method == 'POST':

        form = VehiculoForm(request.POST)

        if form.is_valid():

            form.save()

            return redirect('dashboard')

    else:

        form = VehiculoForm()

    return render(
        request,
        'flota/form_vehiculo.html',
        {
            'form': form,
            'titulo': 'Nuevo vehículo'
        }
    )
@login_required
@editor_required
def editar_vehiculo(request, id):

    vehiculo = Vehiculo.objects.get(
        id=id
    )

    if request.method == 'POST':

        form = VehiculoForm(
            request.POST,
            instance=vehiculo
        )

        if form.is_valid():

            estado_anterior = (
                f"Estado operacional: {vehiculo.estado_operacional} | "
                f"Subestado: {vehiculo.subestado_no_operativo} | "
                f"Kilometraje: {vehiculo.kilometraje_actual}"
            )

            vehiculo_actualizado = form.save()

            estado_nuevo = (
                f"Estado operacional: {vehiculo_actualizado.estado_operacional} | "
                f"Subestado: {vehiculo_actualizado.subestado_no_operativo} | "
                f"Kilometraje: {vehiculo_actualizado.kilometraje_actual}"
            )

            registrar_bitacora(
                request=request,
                accion='EDITAR',
                modulo='Vehículos',
                modelo_afectado='Vehiculo',
                objeto_id=vehiculo_actualizado.id,
                objeto_repr=vehiculo_actualizado.patente,
                descripcion=f'Edición de vehículo {vehiculo_actualizado.patente}.',
                valor_anterior=estado_anterior,
                valor_nuevo=estado_nuevo
            )

            messages.success(
                request,
                f'Vehículo {vehiculo_actualizado.patente} actualizado correctamente.'
            )

            return redirect(
                'detalle_vehiculo',
                id=vehiculo_actualizado.id
            )

    else:

        form = VehiculoForm(
            instance=vehiculo
        )

    return render(

        request,

        'flota/form_vehiculo.html',

        {

            'form':form,

            'titulo':
            f'Editar {vehiculo.patente}'

        }

    )

@login_required
@editor_required
def dar_baja_vehiculo(request, id):

    vehiculo = Vehiculo.objects.get(
        id=id
    )

    if request.method == 'POST':

        motivo_baja = request.POST.get(
            'motivo_baja'
        )

        observacion_baja = request.POST.get(
            'observacion_baja',
            ''
        ).strip()

        if not motivo_baja:

            messages.error(
                request,
                'Debes seleccionar un motivo de baja.'
            )

            return render(
                request,
                'flota/confirmar_baja.html',
                {
                    'vehiculo': vehiculo,
                    'motivos_baja': Vehiculo.MOTIVOS_BAJA,
                    'observacion_baja': observacion_baja,
                }
            )

        estado_anterior = (
            f"Estado administrativo: {vehiculo.estado_administrativo} | "
            f"Estado operacional: {vehiculo.estado_operacional} | "
            f"Subestado: {vehiculo.subestado_no_operativo} | "
            f"Fecha baja: {vehiculo.fecha_baja} | "
            f"Motivo baja: {vehiculo.motivo_baja} | "
            f"Observación baja: {vehiculo.observacion_baja}"
        )

        fecha_baja = timezone.now().date()

        vehiculo.estado_administrativo = 'Dado de Baja'
        vehiculo.estado_operacional = 'No operativo'
        vehiculo.subestado_no_operativo = 'Detenido'
        vehiculo.fecha_baja = fecha_baja
        vehiculo.motivo_baja = motivo_baja
        vehiculo.observacion_baja = observacion_baja

        vehiculo.save()

        HistorialBajaVehiculo.objects.create(
            vehiculo=vehiculo,
            fecha_baja=fecha_baja,
            motivo_baja=motivo_baja,
            observacion_baja=observacion_baja,
            usuario_baja=request.user
        )

        estado_nuevo = (
            f"Estado administrativo: {vehiculo.estado_administrativo} | "
            f"Estado operacional: {vehiculo.estado_operacional} | "
            f"Subestado: {vehiculo.subestado_no_operativo} | "
            f"Fecha baja: {vehiculo.fecha_baja} | "
            f"Motivo baja: {vehiculo.motivo_baja} | "
            f"Observación baja: {vehiculo.observacion_baja}"
        )

        registrar_bitacora(
            request=request,
            accion='BAJA',
            modulo='Vehículos',
            modelo_afectado='Vehiculo',
            objeto_id=vehiculo.id,
            objeto_repr=vehiculo.patente,
            descripcion=(
                f'Vehículo {vehiculo.patente} dado de baja. '
                f'Motivo: {vehiculo.get_motivo_baja_display()}.'
            ),
            valor_anterior=estado_anterior,
            valor_nuevo=estado_nuevo
        )

        messages.success(
            request,
            f'Vehículo {vehiculo.patente} dado de baja correctamente.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    return render(
        request,
        'flota/confirmar_baja.html',
        {
            'vehiculo': vehiculo,
            'motivos_baja': Vehiculo.MOTIVOS_BAJA,
        }
    )


@login_required
@editor_required
def reactivar_vehiculo(request, id):

    vehiculo = Vehiculo.objects.get(
        id=id
    )

    if request.method == 'POST':

        estado_anterior = (
            f"Estado administrativo: {vehiculo.estado_administrativo} | "
            f"Estado operacional: {vehiculo.estado_operacional} | "
            f"Subestado: {vehiculo.subestado_no_operativo} | "
            f"Fecha baja: {vehiculo.fecha_baja} | "
            f"Motivo baja: {vehiculo.motivo_baja} | "
            f"Observación baja: {vehiculo.observacion_baja}"
        )

        historial_abierto = vehiculo.historial_bajas.filter(
            fecha_reactivacion__isnull=True
        ).first()

        if historial_abierto:

            historial_abierto.fecha_reactivacion = timezone.now().date()
            historial_abierto.usuario_reactivacion = request.user
            historial_abierto.save()

        vehiculo.estado_administrativo = 'Vigente'
        vehiculo.estado_operacional = 'No informado'
        vehiculo.subestado_no_operativo = None
        vehiculo.fecha_baja = None
        vehiculo.motivo_baja = None
        vehiculo.observacion_baja = None

        vehiculo.save()

        estado_nuevo = (
            f"Estado administrativo: {vehiculo.estado_administrativo} | "
            f"Estado operacional: {vehiculo.estado_operacional} | "
            f"Subestado: {vehiculo.subestado_no_operativo} | "
            f"Fecha baja: {vehiculo.fecha_baja} | "
            f"Motivo baja: {vehiculo.motivo_baja} | "
            f"Observación baja: {vehiculo.observacion_baja}"
        )

        registrar_bitacora(
            request=request,
            accion='REACTIVAR',
            modulo='Vehículos',
            modelo_afectado='Vehiculo',
            objeto_id=vehiculo.id,
            objeto_repr=vehiculo.patente,
            descripcion=f'Vehículo {vehiculo.patente} reactivado.',
            valor_anterior=estado_anterior,
            valor_nuevo=estado_nuevo
        )

        messages.success(
            request,
            f'{vehiculo.patente} reactivado correctamente.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    return render(
        request,
        'flota/reactivar_vehiculo.html',
        {
            'vehiculo': vehiculo
        }
    )

@login_required
@editor_required
def transferencias_cc(request):

    patente = request.GET.get(
        'patente',
        ''
    ).strip().upper()

    centro_costo_id = request.GET.get(
        'centro_costo',
        ''
    )

    marca = request.GET.get(
        'marca',
        ''
    ).strip().upper()

    vehiculos = Vehiculo.objects.select_related(
        'centro_costo'
    ).exclude(
        estado_administrativo='Dado de Baja'
    ).order_by(
        'patente'
    )

    if patente:

        vehiculos = vehiculos.filter(
            patente__icontains=patente
        )

    if centro_costo_id:

        vehiculos = vehiculos.filter(
            centro_costo_id=centro_costo_id
        )

    if marca:

        ids_marca = []

        for vehiculo in vehiculos.values(
            'id',
            'marca'
        ):

            if normalizar_marca_filtro(
                vehiculo['marca']
            ) == marca:

                ids_marca.append(
                    vehiculo['id']
                )

        vehiculos = vehiculos.filter(
            id__in=ids_marca
        )

    if request.method == 'POST':

        ids = request.POST.getlist(
            'vehiculos'
        )

        centro_destino_id = request.POST.get(
            'centro_costo_destino'
        )

        observacion = request.POST.get(
            'observacion',
            ''
        ).strip()

        if not ids:

            messages.error(
                request,
                'Debes seleccionar al menos un vehículo.'
            )

            return redirect(
                'transferencias_cc'
            )

        if not centro_destino_id:

            messages.error(
                request,
                'Debes seleccionar un centro de costo destino.'
            )

            return redirect(
                'transferencias_cc'
            )

        centro_destino = CentroCosto.objects.get(
            id=centro_destino_id
        )

        vehiculos_a_transferir = Vehiculo.objects.select_related(
            'centro_costo'
        ).filter(
            id__in=ids
        ).exclude(
            estado_administrativo='Dado de Baja'
        )

        transferidos = 0
        omitidos_mismo_cc = 0

        for vehiculo in vehiculos_a_transferir:

            centro_origen = vehiculo.centro_costo

            if (
                centro_origen
                and centro_origen.id == centro_destino.id
            ):

                omitidos_mismo_cc += 1
                continue

            estado_anterior = (
                f"Centro costo anterior: {centro_origen}"
            )

            vehiculo.centro_costo = centro_destino
            vehiculo.save()

            HistorialTransferenciaVehiculo.objects.create(
                vehiculo=vehiculo,
                centro_costo_origen=centro_origen,
                centro_costo_destino=centro_destino,
                fecha_transferencia=timezone.now().date(),
                usuario_transferencia=request.user,
                observacion=observacion
            )

            estado_nuevo = (
                f"Centro costo nuevo: {centro_destino} | "
                f"Observación: {observacion}"
            )

            registrar_bitacora(
                request=request,
                accion='TRANSFERENCIA',
                modulo='Vehículos',
                modelo_afectado='Vehiculo',
                objeto_id=vehiculo.id,
                objeto_repr=vehiculo.patente,
                descripcion=(
                    f'Transferencia masiva de vehículo {vehiculo.patente} '
                    f'de {centro_origen} a {centro_destino}.'
                ),
                valor_anterior=estado_anterior,
                valor_nuevo=estado_nuevo
            )

            transferidos += 1

        if transferidos > 0:

            messages.success(
                request,
                f'{transferidos} vehículos transferidos correctamente.'
            )

        if omitidos_mismo_cc > 0:

            messages.info(
                request,
                f'{omitidos_mismo_cc} vehículos fueron omitidos porque ya pertenecían al centro de costo destino.'
            )

        return redirect(
            'transferencias_cc'
        )

    centros_costo = CentroCosto.objects.order_by(
        'codigo'
    )

    marcas_crudas = Vehiculo.objects.exclude(
        estado_administrativo='Dado de Baja'
    ).exclude(
        marca__isnull=True
    ).exclude(
        marca=''
    ).values_list(
        'marca',
        flat=True
    )

    marcas = sorted(
        {
            normalizar_marca_filtro(
                marca
            )
            for marca in marcas_crudas
            if normalizar_marca_filtro(
                marca
            )
        }
    )

    paginator = Paginator(
        vehiculos,
        25
    )

    page_number = request.GET.get(
        'page'
    )

    vehiculos_pagina = paginator.get_page(
        page_number
    )

    query_params = request.GET.copy()

    if 'page' in query_params:

        query_params.pop(
            'page'
        )

    query_string = query_params.urlencode()

    return render(
        request,
        'flota/transferencias_cc.html',
        {
            'vehiculos': vehiculos_pagina,
            'centros_costo': centros_costo,
            'marcas': marcas,
            'patente_buscada': patente,
            'centro_costo_id': centro_costo_id,
            'marca_seleccionada': marca,
            'query_string': query_string,
        }
    )

@login_required
@editor_required
def transferir_vehiculo(request, id):

    vehiculo = Vehiculo.objects.select_related(
        'centro_costo'
    ).get(
        id=id
    )

    centros_costo = CentroCosto.objects.order_by(
        'codigo'
    )

    if request.method == 'POST':

        centro_destino_id = request.POST.get(
            'centro_costo_destino'
        )

        observacion = request.POST.get(
            'observacion',
            ''
        ).strip()

        if not centro_destino_id:

            messages.error(
                request,
                'Debes seleccionar un centro de costo destino.'
            )

            return render(
                request,
                'flota/transferir_vehiculo.html',
                {
                    'vehiculo': vehiculo,
                    'centros_costo': centros_costo,
                    'observacion': observacion,
                }
            )

        centro_origen = vehiculo.centro_costo

        centro_destino = CentroCosto.objects.get(
            id=centro_destino_id
        )

        if centro_origen and centro_origen.id == centro_destino.id:

            messages.error(
                request,
                'El centro de costo destino debe ser distinto al actual.'
            )

            return render(
                request,
                'flota/transferir_vehiculo.html',
                {
                    'vehiculo': vehiculo,
                    'centros_costo': centros_costo,
                    'observacion': observacion,
                    'centro_destino_id': centro_destino_id,
                }
            )

        estado_anterior = (
            f"Centro costo anterior: {centro_origen}"
        )

        vehiculo.centro_costo = centro_destino
        vehiculo.save()

        HistorialTransferenciaVehiculo.objects.create(
            vehiculo=vehiculo,
            centro_costo_origen=centro_origen,
            centro_costo_destino=centro_destino,
            fecha_transferencia=timezone.now().date(),
            usuario_transferencia=request.user,
            observacion=observacion
        )

        estado_nuevo = (
            f"Centro costo nuevo: {centro_destino} | "
            f"Observación: {observacion}"
        )

        registrar_bitacora(
            request=request,
            accion='TRANSFERENCIA',
            modulo='Vehículos',
            modelo_afectado='Vehiculo',
            objeto_id=vehiculo.id,
            objeto_repr=vehiculo.patente,
            descripcion=(
                f'Vehículo {vehiculo.patente} transferido de '
                f'{centro_origen} a {centro_destino}.'
            ),
            valor_anterior=estado_anterior,
            valor_nuevo=estado_nuevo
        )

        messages.success(
            request,
            f'Vehículo {vehiculo.patente} transferido correctamente.'
        )

        return redirect(
            'detalle_vehiculo',
            id=vehiculo.id
        )

    return render(
        request,
        'flota/transferir_vehiculo.html',
        {
            'vehiculo': vehiculo,
            'centros_costo': centros_costo,
        }
    )
    

@login_required
@master_required
def crear_usuario(request):

    if request.method == 'POST':

        form = CrearUsuarioForm(
            request.POST
        )

        if form.is_valid():

            usuario = form.save()

            rol = form.cleaned_data[
                'rol'
            ]

            usuario.groups.add(
                rol
            )

            usuario.email = (
                form.cleaned_data[
                    'email'
                ]
            )

            usuario.save()

            PerfilUsuario.objects.create(

                usuario=usuario,

                rut=form.cleaned_data[
                    'rut'
                ],

                nombre=form.cleaned_data[
                    'nombre'
                ],

                telefono=form.cleaned_data[
                    'telefono'
                ],

                cargo=form.cleaned_data[
                    'cargo'
                ],

                centro_costo=form.cleaned_data[
                    'centro_costo'
                ],

                sucursal=form.cleaned_data[
                    'sucursal'
                ],
            )

            registrar_bitacora(
                request=request,
                accion='CREAR',
                modulo='Usuarios',
                modelo_afectado='User',
                objeto_id=usuario.id,
                objeto_repr=usuario.username,
                descripcion=f'Usuario {usuario.username} creado.',
                valor_anterior=None,
                valor_nuevo=(
                    f"Usuario: {usuario.username} | "
                    f"Correo: {usuario.email} | "
                    f"Rol: {rol.name} | "
                    f"Nombre: {form.cleaned_data['nombre']} | "
                    f"RUT: {form.cleaned_data['rut']} | "
                    f"Centro costo: {form.cleaned_data['centro_costo']}"
                )
            )

            messages.success(
                request,
                f'Usuario {usuario.username} creado correctamente.'
            )

            return redirect(
                'usuarios'
            )

    else:

        form = CrearUsuarioForm()

    return render(
        request,
        'flota/crear_usuario.html',
        {
            'form': form
        }
    )


@login_required
@master_required
def configuracion(request):

    total_usuarios = User.objects.count()

    return render(
        request,
        'flota/configuracion.html',
        {
            'total_usuarios': total_usuarios
        }
    )
@login_required
@editor_required
def bitacora(request):

    busqueda = request.GET.get(
        'q',
        ''
    ).strip()

    accion = request.GET.get(
        'accion',
        ''
    )

    modulo = request.GET.get(
        'modulo',
        ''
    )

    registros = BitacoraAccion.objects.all()

    if busqueda:

        registros = registros.filter(

            Q(
                usuario_texto__icontains=busqueda
            ) |

            Q(
                descripcion__icontains=busqueda
            ) |

            Q(
                objeto_repr__icontains=busqueda
            ) |

            Q(
                valor_anterior__icontains=busqueda
            ) |

            Q(
                valor_nuevo__icontains=busqueda
            )

        )

    if accion:

        registros = registros.filter(
            accion=accion
        )

    if modulo:

        registros = registros.filter(
            modulo=modulo
        )

    modulos = BitacoraAccion.objects.values_list(
        'modulo',
        flat=True
    ).exclude(
        modulo__isnull=True
    ).distinct().order_by(
        'modulo'
    )

    paginator = Paginator(
        registros,
        25
    )

    page_number = request.GET.get(
        'page'
    )

    registros_pagina = paginator.get_page(
        page_number
    )

    query_params = request.GET.copy()

    if 'page' in query_params:
        query_params.pop(
            'page'
        )

    query_string = query_params.urlencode()

    return render(
        request,
        'flota/bitacora.html',
        {
            'registros': registros_pagina,
            'busqueda': busqueda,
            'accion_seleccionada': accion,
            'modulo_seleccionado': modulo,
            'acciones': BitacoraAccion.ACCIONES,
            'modulos': modulos,
            'query_string': query_string,
        }
    )
    
@login_required
@master_required
def usuarios(request):

    asegurar_perfiles()

    busqueda = request.GET.get(
        'q',
        ''
    )

    usuarios = User.objects.select_related(
        'perfil',
        'perfil__centro_costo'
    ).prefetch_related(
        'groups'
    ).all().order_by(
        'username'
    )

    if busqueda:

        usuarios = usuarios.filter(

            Q(
                username__icontains=busqueda
            ) |

            Q(
                email__icontains=busqueda
            ) |

            Q(
                perfil__nombre__icontains=busqueda
            ) |

            Q(
                perfil__rut__icontains=busqueda
            ) |

            Q(
                groups__name__icontains=busqueda
            )

        ).distinct()

    return render(
        request,
        'flota/usuarios.html',
        {
            'usuarios': usuarios,
            'busqueda': busqueda
        }
    )
    
@login_required
@master_required
def editar_usuario(request, id):

    usuario = User.objects.get(
        id=id
    )

    perfil, _ = PerfilUsuario.objects.get_or_create(
        usuario=usuario
    )

    if request.method == 'POST':

        grupos_anteriores = ", ".join(
            usuario.groups.values_list(
                'name',
                flat=True
            )
        )

        estado_anterior = (
            f"Usuario: {usuario.username} | "
            f"Correo: {usuario.email} | "
            f"Nombre: {perfil.nombre} | "
            f"RUT: {perfil.rut} | "
            f"Teléfono: {perfil.telefono} | "
            f"Cargo: {perfil.cargo} | "
            f"Centro costo: {perfil.centro_costo} | "
            f"Sucursal: {perfil.sucursal} | "
            f"Roles: {grupos_anteriores}"
        )

        usuario.email = request.POST.get(
            'email'
        )

        usuario.save()

        perfil.rut = request.POST.get(
            'rut'
        )

        perfil.nombre = request.POST.get(
            'nombre'
        )

        perfil.telefono = request.POST.get(
            'telefono'
        )

        perfil.cargo = request.POST.get(
            'cargo'
        )

        perfil.sucursal = request.POST.get(
            'sucursal'
        )

        centro_id = request.POST.get(
            'centro_costo'
        )

        if centro_id:

            perfil.centro_costo = CentroCosto.objects.get(
                id=centro_id
            )

        else:

            perfil.centro_costo = None

        perfil.save()

        grupo_id = request.POST.get(
            'rol'
        )

        usuario.groups.clear()

        grupo = Group.objects.get(
            id=grupo_id
        )

        usuario.groups.add(
            grupo
        )

        grupos_nuevos = ", ".join(
            usuario.groups.values_list(
                'name',
                flat=True
            )
        )

        estado_nuevo = (
            f"Usuario: {usuario.username} | "
            f"Correo: {usuario.email} | "
            f"Nombre: {perfil.nombre} | "
            f"RUT: {perfil.rut} | "
            f"Teléfono: {perfil.telefono} | "
            f"Cargo: {perfil.cargo} | "
            f"Centro costo: {perfil.centro_costo} | "
            f"Sucursal: {perfil.sucursal} | "
            f"Roles: {grupos_nuevos}"
        )

        registrar_bitacora(
            request=request,
            accion='EDITAR',
            modulo='Usuarios',
            modelo_afectado='User',
            objeto_id=usuario.id,
            objeto_repr=usuario.username,
            descripcion=f'Usuario {usuario.username} actualizado.',
            valor_anterior=estado_anterior,
            valor_nuevo=estado_nuevo
        )

        messages.success(
            request,
            "Usuario actualizado."
        )

        return redirect(
            'usuarios'
        )

    grupos = Group.objects.filter(
        name__in=[
            'Visualizador',
            'Editor',
            'Master'
        ]
    )

    centros = CentroCosto.objects.order_by(
        'codigo'
    )

    return render(
        request,
        'flota/editar_usuario.html',
        {
            'usuario_obj': usuario,
            'perfil': perfil,
            'grupos': grupos,
            'centros': centros
        }
    )

def asegurar_perfiles():

    usuarios = User.objects.all()

    for usuario in usuarios:

        PerfilUsuario.objects.get_or_create(
            usuario=usuario
        )

@login_required
def eliminar_usuario(
    request,
    id
):

    if not request.user.groups.filter(
        name='Master'
    ).exists():

        return redirect(
            'dashboard'
        )

    usuario = get_object_or_404(
        User,
        id=id
    )

    if request.method == 'POST':

        if usuario == request.user:

            messages.error(
                request,
                'No puedes eliminar tu propio usuario.'
            )

            return redirect(
                'usuarios'
            )

        estado_anterior = (
            f"Usuario: {usuario.username} | "
            f"Correo: {usuario.email} | "
            f"Nombre: {getattr(usuario.perfil, 'nombre', None)} | "
            f"RUT: {getattr(usuario.perfil, 'rut', None)} | "
            f"Centro costo: {getattr(usuario.perfil, 'centro_costo', None)} | "
            f"Roles: {', '.join(usuario.groups.values_list('name', flat=True))}"
        )

        usuario_id = usuario.id
        usuario_username = usuario.username

        usuario.delete()

        registrar_bitacora(
            request=request,
            accion='ELIMINAR',
            modulo='Usuarios',
            modelo_afectado='User',
            objeto_id=usuario_id,
            objeto_repr=usuario_username,
            descripcion=f'Usuario {usuario_username} eliminado.',
            valor_anterior=estado_anterior,
            valor_nuevo=None
        )

        messages.success(
            request,
            'Usuario eliminado correctamente.'
        )

        return redirect(
            'usuarios'
        )

    return render(
        request,
        'flota/confirmar_eliminar_usuario.html',
        {
            'usuario': usuario
        }
    )

@login_required
@master_required
def resetear_password_usuario(request, id):

    usuario = get_object_or_404(
        User,
        id=id
    )

    if request.method == 'POST':

        form = SetPasswordForm(
            usuario,
            request.POST
        )

        if form.is_valid():

            form.save()

            registrar_bitacora(
                request=request,
                accion='RESET_PASSWORD',
                modulo='Usuarios',
                modelo_afectado='User',
                objeto_id=usuario.id,
                objeto_repr=usuario.username,
                descripcion=(
                    f'Contraseña del usuario {usuario.username} '
                    f'reseteada por administrador.'
                ),
                valor_anterior=None,
                valor_nuevo='Contraseña actualizada. No se registra el valor por seguridad.'
            )

            messages.success(
                request,
                f'Contraseña de {usuario.username} actualizada correctamente.'
            )

            return redirect(
                'usuarios'
            )

    else:

        form = SetPasswordForm(
            usuario
        )

    return render(
        request,
        'flota/resetear_password_usuario.html',
        {
            'form': form,
            'usuario_obj': usuario
        }
    )


@login_required
@editor_required
def registrar_mantencion_vehiculo(request, id):

    vehiculo = Vehiculo.objects.get(
        id=id
    )

    kilometraje_programado_sugerido = obtener_kilometraje_programado_mantencion(
        vehiculo
    )

    mantencion_en_curso = obtener_mantencion_en_curso_vehiculo(
        vehiculo
    )

    if request.method == 'POST':

        datos_post = request.POST.copy()

        if (
            datos_post.get('tipo_mantencion') == 'KILOMETRAJE'
            and not datos_post.get('kilometraje_programado')
            and kilometraje_programado_sugerido
        ):

            datos_post['kilometraje_programado'] = str(
                kilometraje_programado_sugerido
            )

        form = MantencionVehiculoForm(
            datos_post
        )

        if form.is_valid():

            estado_solicitado = form.cleaned_data.get(
                'estado'
            )

            if (
                estado_solicitado == 'EN_CURSO'
                and mantencion_en_curso
            ):

                messages.error(
                    request,
                    (
                        f'No se puede registrar una nueva mantención en curso '
                        f'para {vehiculo.patente}, porque ya existe una '
                        f'mantención activa. Debes cerrar la mantención actual '
                        f'antes de iniciar otra.'
                    )
                )

                return render(
                    request,
                    'flota/registrar_mantencion.html',
                    {
                        'form': form,
                        'vehiculo': vehiculo,
                        'kilometraje_programado_sugerido': kilometraje_programado_sugerido,
                        'mantencion_en_curso': mantencion_en_curso,
                    }
                )

            mantencion = form.save(
                commit=False
            )

            mantencion.vehiculo = vehiculo
            mantencion.usuario_registro = request.user

            if (
                mantencion.estado == 'EN_CURSO'
                and not mantencion.fecha_ingreso
            ):

                mantencion.fecha_ingreso = timezone.now().date()

            if (
                mantencion.estado == 'EN_CURSO'
                and not mantencion.kilometraje_ingreso
                and vehiculo.kilometraje_actual
            ):

                mantencion.kilometraje_ingreso = vehiculo.kilometraje_actual

            mantencion.save()

            estado_vehiculo_anterior = (
                f"Estado operacional: {vehiculo.estado_operacional} | "
                f"Subestado: {vehiculo.subestado_no_operativo}"
            )

            vehiculo_actualizado_por_mantencion = False

            if mantencion.estado == 'EN_CURSO':

                vehiculo.estado_operacional = 'No operativo'

                if mantencion.tipo_mantencion in [
                    'PROGRAMADA',
                    'KILOMETRAJE',
                ]:

                    vehiculo.subestado_no_operativo = 'Mantencion'

                elif mantencion.tipo_mantencion == 'SINIESTRO':

                    vehiculo.subestado_no_operativo = 'Reparacion'

                vehiculo.save()

                vehiculo_actualizado_por_mantencion = True

                estado_vehiculo_nuevo = (
                    f"Estado operacional: {vehiculo.estado_operacional} | "
                    f"Subestado: {vehiculo.subestado_no_operativo}"
                )

                registrar_bitacora(
                    request=request,
                    accion='CAMBIO_ESTADO',
                    modulo='Vehículos',
                    modelo_afectado='Vehiculo',
                    objeto_id=vehiculo.id,
                    objeto_repr=vehiculo.patente,
                    descripcion=(
                        f'Vehículo {vehiculo.patente} actualizado automáticamente '
                        f'por inicio de mantención.'
                    ),
                    valor_anterior=estado_vehiculo_anterior,
                    valor_nuevo=estado_vehiculo_nuevo
                )

            registrar_bitacora(
                request=request,
                accion='CREAR',
                modulo='Mantenciones',
                modelo_afectado='MantencionVehiculo',
                objeto_id=mantencion.id,
                objeto_repr=(
                    f'{vehiculo.patente} - '
                    f'{mantencion.get_tipo_mantencion_display()}'
                ),
                descripcion=(
                    f'Registro de mantención para vehículo {vehiculo.patente}. '
                    f'Tipo: {mantencion.get_tipo_mantencion_display()}. '
                    f'Estado: {mantencion.get_estado_display()}.'
                ),
                valor_anterior=None,
                valor_nuevo=(
                    f"Vehículo: {vehiculo.patente} | "
                    f"Tipo: {mantencion.get_tipo_mantencion_display()} | "
                    f"Estado: {mantencion.get_estado_display()} | "
                    f"Fecha programada: {mantencion.fecha_programada} | "
                    f"Kilometraje programado: {mantencion.kilometraje_programado} | "
                    f"Fecha ingreso: {mantencion.fecha_ingreso} | "
                    f"Kilometraje ingreso: {mantencion.kilometraje_ingreso} | "
                    f"Motivo: {mantencion.motivo} | "
                    f"Observación: {mantencion.observacion}"
                )
            )

            if vehiculo_actualizado_por_mantencion:

                if mantencion.tipo_mantencion == 'SINIESTRO':

                    messages.success(
                        request,
                        (
                            f'Mantención por siniestro registrada correctamente '
                            f'para {vehiculo.patente}. '
                            f'El vehículo fue actualizado automáticamente a '
                            f'No operativo / Reparación.'
                        )
                    )

                else:

                    messages.success(
                        request,
                        (
                            f'Mantención registrada correctamente para '
                            f'{vehiculo.patente}. '
                            f'El vehículo fue actualizado automáticamente a '
                            f'No operativo / Mantención.'
                        )
                    )

            else:

                messages.success(
                    request,
                    (
                        f'Mantención registrada correctamente para '
                        f'{vehiculo.patente}. '
                        f'El estado operacional del vehículo no fue modificado '
                        f'porque la mantención aún no está en curso.'
                    )
                )

            return redirect(
                'detalle_vehiculo',
                id=vehiculo.id
            )

    else:

        valores_iniciales = {}

        if kilometraje_programado_sugerido:

            valores_iniciales['kilometraje_programado'] = kilometraje_programado_sugerido

        form = MantencionVehiculoForm(
            initial=valores_iniciales
        )

    return render(
        request,
        'flota/registrar_mantencion.html',
        {
            'form': form,
            'vehiculo': vehiculo,
            'kilometraje_programado_sugerido': kilometraje_programado_sugerido,
            'mantencion_en_curso': mantencion_en_curso,
        }
    )
 
@login_required
def control_mantenciones(request):

    hoy = timezone.now().date()

    fecha_limite_proximas = hoy + timedelta(
        days=15
    )

    resumen_mantenciones_km = obtener_resumen_mantenciones_km()

    alertas_km_vencidas = resumen_mantenciones_km[
        'alertas_km_vencidas'
    ]

    alertas_km_proximas = resumen_mantenciones_km[
        'alertas_km_proximas'
    ]

    alertas_km_al_dia = resumen_mantenciones_km[
        'alertas_km_al_dia'
    ]

    alertas_km_sin_datos = resumen_mantenciones_km[
        'alertas_km_sin_datos'
    ]

    mantenciones_programadas = MantencionVehiculo.objects.select_related(
        'vehiculo',
        'vehiculo__centro_costo',
        'usuario_registro'
    ).filter(
        estado='PENDIENTE',
        tipo_mantencion='PROGRAMADA',
        fecha_programada__isnull=False
    ).exclude(
        vehiculo__estado_administrativo='Dado de Baja'
    ).order_by(
        'fecha_programada'
    )

    mantenciones_programadas_vencidas = mantenciones_programadas.filter(
        fecha_programada__lt=hoy
    )

    mantenciones_programadas_proximas = mantenciones_programadas.filter(
        fecha_programada__gte=hoy,
        fecha_programada__lte=fecha_limite_proximas
    )

    mantenciones_en_curso = MantencionVehiculo.objects.select_related(
        'vehiculo',
        'vehiculo__centro_costo',
        'usuario_registro'
    ).filter(
        estado='EN_CURSO'
    ).exclude(
        vehiculo__estado_administrativo='Dado de Baja'
    ).order_by(
        'fecha_ingreso',
        'fecha_programada',
        'vehiculo__patente'
    )

    vehiculos_vigentes_operacionales = Vehiculo.objects.exclude(
        estado_administrativo='Dado de Baja'
    )

    total_vehiculos_vigentes = vehiculos_vigentes_operacionales.count()

    total_vehiculos_operativos = vehiculos_vigentes_operacionales.filter(
        estado_operacional='Operativo'
    ).count()

    total_vehiculos_no_operativos = vehiculos_vigentes_operacionales.filter(
        estado_operacional='No operativo'
    ).count()

    total_vehiculos_en_mantencion = vehiculos_vigentes_operacionales.filter(
        estado_operacional='No operativo',
        subestado_no_operativo='Mantencion'
    ).count()

    total_vehiculos_en_reparacion = vehiculos_vigentes_operacionales.filter(
        estado_operacional='No operativo',
        subestado_no_operativo='Reparacion'
    ).count()

    total_vehiculos_detenidos = vehiculos_vigentes_operacionales.filter(
        estado_operacional='No operativo',
        subestado_no_operativo='Detenido'
    ).count()

    if total_vehiculos_vigentes:

        porcentaje_vehiculos_operativos = round(
            (
                total_vehiculos_operativos
                / total_vehiculos_vigentes
            ) * 100,
            1
        )

        porcentaje_vehiculos_no_operativos = round(
            (
                total_vehiculos_no_operativos
                / total_vehiculos_vigentes
            ) * 100,
            1
        )

    else:

        porcentaje_vehiculos_operativos = 0
        porcentaje_vehiculos_no_operativos = 0
    
    contexto = {
        'hoy': hoy,

        'alertas_km_vencidas':
        alertas_km_vencidas,

        'alertas_km_proximas':
        alertas_km_proximas,

        'alertas_km_al_dia':
        alertas_km_al_dia,

        'alertas_km_sin_datos':
        alertas_km_sin_datos,

        'mantenciones_programadas_vencidas':
        mantenciones_programadas_vencidas,

        'mantenciones_programadas_proximas':
        mantenciones_programadas_proximas,

        'mantenciones_en_curso':
        mantenciones_en_curso,

        'total_programadas_vencidas':
        mantenciones_programadas_vencidas.count(),

        'total_programadas_proximas':
        mantenciones_programadas_proximas.count(),

        'total_en_curso':
        mantenciones_en_curso.count(),

        'total_vehiculos_vigentes':
        total_vehiculos_vigentes,

        'total_vehiculos_operativos':
        total_vehiculos_operativos,

        'total_vehiculos_no_operativos':
        total_vehiculos_no_operativos,

        'total_vehiculos_en_mantencion':
        total_vehiculos_en_mantencion,

        'total_vehiculos_en_reparacion':
        total_vehiculos_en_reparacion,

        'total_vehiculos_detenidos':
        total_vehiculos_detenidos,

        'porcentaje_vehiculos_operativos':
        porcentaje_vehiculos_operativos,

        'porcentaje_vehiculos_no_operativos':
        porcentaje_vehiculos_no_operativos,
        
        'puede_editar': (
            request.user.is_superuser
            or request.user.groups.filter(name='Editor').exists()
            or request.user.groups.filter(name='Master').exists()
        ),
    }

    contexto.update(
        resumen_mantenciones_km
    )

    return render(
        request,
        'flota/control_mantenciones.html',
        contexto
    )
    
@login_required
def exportar_control_mantenciones_excel(request):

    secciones_validas = [
        'resumen',
        'km_vencidas',
        'km_proximas',
        'sin_datos',
        'al_dia',
        'programadas_vencidas',
        'programadas_proximas',
        'en_curso',
    ]

    secciones_exportar = request.GET.getlist(
        'secciones'
    )

    if not secciones_exportar:

        secciones_exportar = secciones_validas

    secciones_exportar = [
        seccion for seccion in secciones_exportar
        if seccion in secciones_validas
    ]

    if not secciones_exportar:

        secciones_exportar = [
            'resumen'
        ]

    hoy = timezone.now().date()

    fecha_limite_proximas = hoy + timedelta(
        days=15
    )

    resumen_mantenciones_km = obtener_resumen_mantenciones_km()

    alertas_km_vencidas = resumen_mantenciones_km[
        'alertas_km_vencidas'
    ]

    alertas_km_proximas = resumen_mantenciones_km[
        'alertas_km_proximas'
    ]

    alertas_km_sin_datos = resumen_mantenciones_km[
        'alertas_km_sin_datos'
    ]

    alertas_km_al_dia = resumen_mantenciones_km[
        'alertas_km_al_dia'
    ]

    mantenciones_programadas = MantencionVehiculo.objects.select_related(
        'vehiculo',
        'vehiculo__centro_costo',
        'usuario_registro'
    ).filter(
        estado='PENDIENTE',
        tipo_mantencion='PROGRAMADA',
        fecha_programada__isnull=False
    ).exclude(
        vehiculo__estado_administrativo='Dado de Baja'
    ).order_by(
        'fecha_programada'
    )

    mantenciones_programadas_vencidas = mantenciones_programadas.filter(
        fecha_programada__lt=hoy
    )

    mantenciones_programadas_proximas = mantenciones_programadas.filter(
        fecha_programada__gte=hoy,
        fecha_programada__lte=fecha_limite_proximas
    )

    mantenciones_en_curso = MantencionVehiculo.objects.select_related(
        'vehiculo',
        'vehiculo__centro_costo',
        'usuario_registro'
    ).filter(
        estado='EN_CURSO'
    ).exclude(
        vehiculo__estado_administrativo='Dado de Baja'
    ).order_by(
        'fecha_ingreso',
        'fecha_programada',
        'vehiculo__patente'
    )
        
    wb = Workbook()

    hoja_inicial = wb.active
    wb.remove(
        hoja_inicial
    )

    def ajustar_columnas(hoja):

        for columna in hoja.columns:

            max_largo = 0
            letra_columna = columna[0].column_letter

            for celda in columna:

                if celda.value is not None:

                    max_largo = max(
                        max_largo,
                        len(
                            str(
                                celda.value
                            )
                        )
                    )

            hoja.column_dimensions[letra_columna].width = min(
                max_largo + 2,
                45
            )

    def agregar_hoja_resumen():

        ws = wb.create_sheet(
            title='Resumen'
        )

        ws.append([
            'Indicador',
            'Cantidad',
            'Porcentaje'
        ])

        ws.append([
            'Vehículos controlados',
            resumen_mantenciones_km['total_vehiculos_controlados'],
            '100%'
        ])

        ws.append([
            'Al día',
            resumen_mantenciones_km['total_km_al_dia'],
            f"{resumen_mantenciones_km['porcentaje_km_al_dia']}%"
        ])

        ws.append([
            'Próximas por KM',
            resumen_mantenciones_km['total_km_proximas'],
            f"{resumen_mantenciones_km['porcentaje_km_proximas']}%"
        ])

        ws.append([
            'Vencidas por KM',
            resumen_mantenciones_km['total_km_vencidas'],
            f"{resumen_mantenciones_km['porcentaje_km_vencidas']}%"
        ])

        ws.append([
            'Sin información',
            resumen_mantenciones_km['total_km_sin_datos'],
            f"{resumen_mantenciones_km['porcentaje_km_sin_datos']}%"
        ])

        ws.append([
            'Programadas vencidas',
            mantenciones_programadas_vencidas.count(),
            ''
        ])

        ws.append([
            'Programadas próximas',
            mantenciones_programadas_proximas.count(),
            ''
        ])

        ws.append([
            'Mantenciones en curso',
            mantenciones_en_curso.count(),
            ''
        ])

        ajustar_columnas(
            ws
        )

    def agregar_hoja_alertas_km(
        nombre_hoja,
        alertas
    ):

        ws = wb.create_sheet(
            title=nombre_hoja
        )

        ws.append([
            'Patente',
            'Centro de costo',
            'Kilometraje actual',
            'Kilometraje mantención',
            'Kilómetros restantes',
            'Estado alerta'
        ])

        for alerta in alertas:

            vehiculo = alerta['vehiculo']

            ws.append([
                vehiculo.patente,
                str(
                    vehiculo.centro_costo
                ) if vehiculo.centro_costo else '',
                alerta.get(
                    'kilometraje_actual'
                ),
                alerta.get(
                    'kilometraje_programado'
                ),
                alerta.get(
                    'kilometros_restantes'
                ),
                alerta.get(
                    'estado_alerta'
                ),
            ])

        ajustar_columnas(
            ws
        )

    def agregar_hoja_sin_datos(
        nombre_hoja,
        alertas
    ):

        ws = wb.create_sheet(
            title=nombre_hoja
        )

        ws.append([
            'Patente',
            'Centro de costo',
            'Kilometraje actual',
            'Última mantención',
            'Próxima mantención',
            'Problemas detectados'
        ])

        for alerta in alertas:

            vehiculo = alerta['vehiculo']

            problemas = ', '.join(
                alerta.get(
                    'problemas',
                    []
                )
            )

            ws.append([
                vehiculo.patente,
                str(
                    vehiculo.centro_costo
                ) if vehiculo.centro_costo else '',
                vehiculo.kilometraje_actual,
                vehiculo.kilometraje_ultima_mantencion,
                vehiculo.kilometraje_proxima_mantencion,
                problemas,
            ])

        ajustar_columnas(
            ws
        )

    def agregar_hoja_mantenciones(
        nombre_hoja,
        mantenciones
    ):

        ws = wb.create_sheet(
            title=nombre_hoja
        )

        ws.append([
            'Patente',
            'Centro de costo',
            'Tipo',
            'Estado',
            'Fecha programada',
            'Fecha ingreso',
            'Kilometraje programado',
            'Kilometraje ingreso',
            'Motivo',
            'Usuario registro'
        ])

        for mantencion in mantenciones:

            vehiculo = mantencion.vehiculo

            ws.append([
                vehiculo.patente,
                str(
                    vehiculo.centro_costo
                ) if vehiculo.centro_costo else '',
                mantencion.get_tipo_mantencion_display(),
                mantencion.get_estado_display(),
                mantencion.fecha_programada,
                mantencion.fecha_ingreso,
                mantencion.kilometraje_programado,
                mantencion.kilometraje_ingreso,
                mantencion.motivo,
                mantencion.usuario_registro.username
                if mantencion.usuario_registro
                else '',
            ])

        ajustar_columnas(
            ws
        )

    if 'resumen' in secciones_exportar:

        agregar_hoja_resumen()

    if 'km_vencidas' in secciones_exportar:

        agregar_hoja_alertas_km(
            'KM vencidas',
            alertas_km_vencidas
        )

    if 'km_proximas' in secciones_exportar:

        agregar_hoja_alertas_km(
            'KM proximas',
            alertas_km_proximas
        )

    if 'sin_datos' in secciones_exportar:

        agregar_hoja_sin_datos(
            'Sin informacion',
            alertas_km_sin_datos
        )

    if 'al_dia' in secciones_exportar:

        agregar_hoja_alertas_km(
            'Al dia',
            alertas_km_al_dia
        )

    if 'programadas_vencidas' in secciones_exportar:

        agregar_hoja_mantenciones(
            'Programadas vencidas',
            mantenciones_programadas_vencidas
        )

    if 'programadas_proximas' in secciones_exportar:

        agregar_hoja_mantenciones(
            'Programadas proximas',
            mantenciones_programadas_proximas
        )

    if 'en_curso' in secciones_exportar:

        agregar_hoja_mantenciones(
            'En curso',
            mantenciones_en_curso
        )

    if not wb.worksheets:

        agregar_hoja_resumen()

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="control_mantenciones.xlsx"'
    )

    wb.save(
        response
    )

    return response
